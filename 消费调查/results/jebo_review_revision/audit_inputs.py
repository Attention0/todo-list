import argparse, re, json, hashlib
from pathlib import Path
from docx import Document
from common import *
p=argparse.ArgumentParser();p.add_argument('--data',required=True);p.add_argument('--questionnaire',required=True);a=p.parse_args()
doc=Document(a.questionnaire);parts=[p.text for p in doc.paragraphs]
for t in doc.tables:
    parts += [' | '.join(c.text for c in r.cells) for r in t.rows]
text='\n'.join(parts); start=text.index('随机情景题');end=text.find('\n附',start); scenarios=text[start:end if end>start else len(text)]
sections=re.split(r'(?=版本 \d+（数据题号)',scenarios)[1:];assert len(sections)==9
rows=[]
for j,s in enumerate(sections):
    lines=[x.strip() for x in s.splitlines() if x.strip()]; opts=[x for x in lines if re.match('[①②③④⑤⑥]',x)];assert len(opts)==6
    for k,line in enumerate(opts,1):
        rows.append(dict(form=FORMS[j//3],amount=int(AMOUNTS[j%3]),category=k,chinese=line,
         english=['Essentially no additional consumption; parenthetical saving/debt explanation differs by form','A small part, approximately below10%; yuan example from0 to0.10T','A part, approximately10–25%; yuan example0.10T–0.25T','About half, approximately25–50%; yuan example0.25T–0.50T','Most, approximately50–75%; yuan example0.50T–0.75T','Almost all, approximately above75%; yuan example0.75T–T'][k-1],
         exact_lower='No verbal category1 cutoff' if k==1 else 'Approximate yuan example, not exact spending',
         exact_upper='Absent' if k==1 else 'Top verbal threshold open; parenthetical example ends at T' if k==6 else 'Approximate yuan example'))
save(rows,'questionnaire_options.csv')
note('QUESTION_INTERVAL_AUDIT.md',f'''# Original questionnaire interval audit

Original integrated Chinese DOCX SHA256: {hashlib.sha256(Path(a.questionnaire).read_bytes()).hexdigest()}.
Read all nine versions, not just the codebook. All54 options and translations are
recorded in questionnaire_options.csv and the appendix.

Category1 is verbal “基本不会额外多消费”, with form-specific saving/debt explanations.
It is not an exact observed zero and has no numeric boundary with category2.
Category2 says “约10%以下” with a0–10% yuan example. The interior percentages and
examples are approximate. Category6 says “约75%以上” but its example ends at100%.
The verbal threshold is open; the example is not an explicit prohibition on spending
more than the transfer. No numerical separation of the bottom two bins or100% cap
will be invented.

Gate: proceed only as a sensitivity mapping. Merging categories1+2 below10% is
compatible with the ordinary meaning of “essentially no additional consumption”
and the second category's example, conditional on a respondent interpreting the
bottom verbal label below the10% threshold. This implication is not an exact
measurement fact. The five merged intervals use left censoring at.10T, then.10–.25T,
.25–.50T,.50–.75T, and right censoring at.75T. The latent normal model may permit
negative spending; its left-censored tail is not an observation of negative or zero
actual expenditure. Approximate wording and implicit horizons remain limitations.
Original six-category counts, ordinal and midpoint results remain primary observed
descriptions. No interval model creates coding-free identification.

The questionnaire describes system random selection of one of nine versions, and
one answered scenario per record. It does not document implementation probabilities,
random-number generator, allocation logs or field dates. Food has six-month expiry;
Medical is described as long-term. There is no common explicit total-spending horizon.
''')
app=['# Exact questionnaire appendix','', 'Chinese text is transcribed verbatim from the original integrated questionnaire. English is a translation, not a separately administered instrument. Wording and approximate numeric examples are retained.','']
app.extend(['## Randomization and common instruction','',scenarios[:scenarios.index('版本 1')].strip(),''])
for j,s in enumerate(sections):
    lines=[x.strip() for x in s.splitlines() if x.strip()]
    app += ['## '+lines[0],'', '\n\n'.join(lines[1:]),'', 'English scenario: '+(['A hypothetical one-time government cash transfer, deposited into a bank account or WeChat/Alipay, with unrestricted use for spending, saving or debt.','A hypothetical electronic voucher limited to food and daily necessities at supermarkets, markets and convenience stores; valid six months; noncashable and unavailable for other uses.','A hypothetical credit to the medical-insurance personal account for eligible medical spending, such as visits, medicines and checks; long-term validity.'][j//3])+f' Amount RMB{AMOUNTS[j%3]:,}. Outcome: by how much would your TOTAL consumption exceed your original plan?', '']
    for r in rows[j*6:j*6+6]:app.append(str(r['category'])+'. '+r['english'])
    app.append('')
note('QUESTIONNAIRE_APPENDIX.md','\n'.join(app))
reader=pd.io.stata.StataReader(a.data,convert_categoricals=False)
labels=reader.variable_labels();values=reader.value_labels()
for key in ['q28_hhsize','q29_foodexp','q41_gender','q43_edu','q46_income','q24_workstat','q49_citytier']:
    print(key,labels.get(key),next((v for n,v in values.items() if n==key),{}))
# Read only study-relevant document text for provenance; manuscript assertions do
# not substitute for ethics/recruitment records.
hits=[]
for path in Path(a.questionnaire).parent.glob('*.docx'):
    if path.name==Path(a.questionnaire).name:continue
    try:
        for p in Document(path).paragraphs:
            if re.search(r'ethic|consent|recruit|field(?:work| date)|腾讯|202[0-9]年|知情|伦理|funding|conflict|compensat',p.text,re.I):
                hits.append(dict(source=path.name,text=p.text[:1600]))
    except Exception as e:hits.append(dict(source=path.name,text='Read failure: '+type(e).__name__))
save(hits or [dict(source='No matched record',text='')],'metadata_document_hits.csv')
note('STUDY_METADATA_AUDIT.md','''# Study metadata audit

Searched the repository audit/manuscript records, original integrated questionnaire,
authorized data labels and all manuscript DOCX files in the consumption-study source
folder. Assertions extracted from drafts are catalogued in metadata_document_hits.csv;
they are not independently verified collection records. Sources/ and local manuscripts
were read only.

Verified from the instrument/data: one of nine scenarios is selected by the system;
each record answers one scenario; platform-background fields follow the numbered
questionnaire; adult exclusion is recorded age below18; six original response categories.
Earlier platform identification is Tencent Questionnaire; it does not establish the
recruitment vendor, sampling frame or fieldwork dates. Prior author names/affiliations
are retained for author verification. Dataset delivery contains5,497 records,17 recorded
minors; adult analysis does not repair collection ethics.

No separate approval/exemption certificate, consent form, invitation/start/completion
flow, allocation implementation/log, field dates, incentive record or stopping protocol
was found in the available project sources. Recruitment claims in local drafts remain
unverified. Funding, conflict, CRediT, sharing permissions and authors' AI-disclosure
confirmation require author input. The concrete questions appear in the author checklist.
''')
old={}
for parent in [ROOT/'results',ROOT/'manuscript']:
    for path in parent.rglob('*'):
        if path.is_file() and not any(x in path.parts for x in ['jebo_review_revision','jebo_v3','qa','research_cache','__pycache__','logs']):
            old[str(path.relative_to(ROOT)).replace('\\','/')]=hashlib.sha256(path.read_bytes()).hexdigest()
(O/'historical_hashes.json').write_text(json.dumps(old,indent=2,ensure_ascii=False)+'\n',encoding='utf8')
print('Questionnaire gate published; historical files hashed:',len(old))
