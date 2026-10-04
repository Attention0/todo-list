from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
O=Path(__file__).resolve().parent;M=O.parents[1]/'manuscript'/'jebo_v3'
def text(p,t):
    for j,x in enumerate(re.split(r'\*\*(.*?)\*\*',t)):
        r=p.add_run(x);r.bold=bool(j%2)
def newdoc():
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.78);sec.left_margin=sec.right_margin=Inches(.82)
    for st in d.styles:
        if st.type==1:
            st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
            rp=st.element.get_or_add_rPr();fonts=rp.find(qn('w:rFonts'))
            if fonts is not None:fonts.set(qn('w:eastAsia'),'SimSun')
    st=d.styles['Normal'];st.font.size=Pt(11.5);st.paragraph_format.line_spacing=1.2;st.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',19),('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.1
    d.styles['Title'].paragraph_format.space_before=Pt(0)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    d.core_properties.author='';d.core_properties.subject='Reviewer revision draft of stated transfer spending experiment'
    return d
def tab(d,b):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in b.splitlines() if not re.match(r'^\|[\s:|\-]+\|$',line)]
    n=len(rows[0]);assert all(len(r)==n for r in rows),(n,rows)
    t=d.add_table(rows=0,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    if n==4:w=[1.7,1.64,1.64,1.65]
    elif n==5:w=[.72,1.30,1.30,1.65,1.66]
    else:
        first=1.2 if n>6 else 1.6
        w=[first]+[(6.63-first)/(n-1)]*(n-1)
    if rows[0][0]=='account':w=[2.05,.43]+[(6.63-2.48)/(n-2)]*(n-2)
    if rows[0][0]=='model' and n>=7:w=[1.16,1.28]+[(6.63-2.44)/(n-2)]*(n-2)
    if rows[0][0]=='comparison':w=[2.35]+[(6.63-2.35)/(n-1)]*(n-1)
    if rows[0][:3]==['variable','category','label']:w=[1.01,.56,2.02]+[(6.63-3.59)/(n-3)]*(n-3)
    if n==9:w=[6.63/n]*n
    if n>=10:w=[6.63/n]*n
    for col,wide in zip(t.columns,w):col.width=Inches(wide)
    for j,row in enumerate(rows):
        cells=t.add_row().cells;pr=cells[0]._tc.getparent().get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if j==0:pr.append(OxmlElement('w:tblHeader'))
        for i,(c,v) in enumerate(zip(cells,row)):
            c.width=Inches(w[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=v.replace('_',' ') if j==0 else v
            props=c._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'60');el.set(qn('w:type'),'dxa');margin.append(el)
            props.append(margin);borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            props.append(borders)
            if j==0:shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFEFEF');props.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.keep_with_next=j==0 or (len(rows)<10 and j<len(rows)-1)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 or (rows[0][i] in ['label','parameter','comparison','type']) else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9.5);r.bold=j==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)
def build(stem):
    d=newdoc();src=(M/(stem+'.md')).read_text(encoding='utf8');ref=False;appendix=False
    for b in src.split('\n\n'):
        b=b.strip()
        if not b:continue
        if b.startswith('# '):
            title=b[2:];d.add_paragraph(title,'Title');d.core_properties.title=title
        elif b.startswith('## '):
            name=b[3:];ref=name=='References';appendix=name=='Tables and figures'
            if appendix:continue
            p=d.add_paragraph(name,'Heading 1')
        elif b.startswith('### '):
            p=d.add_paragraph(b[4:],'Heading 2')
            if appendix and re.match(r'Table [13]\.',b[4:]):p.paragraph_format.page_break_before=True
        elif b.startswith('|'):tab(d,b)
        elif b.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',b);p=d.add_paragraph();p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True;r=p.add_run();r.add_picture(str(M/m.group(2)),width=Inches(6.6));r._r.xpath('.//wp:docPr')[0].set('descr',m.group(1))
        else:
            p=d.add_paragraph();text(p,b)
            if ref:
                p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(10)
            if re.match(r'^Figure \d+\.',b):
                p.paragraph_format.line_spacing=1.1
                for r in p.runs:r.font.size=Pt(10)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.save(M/(stem+'.docx'));return dict(file=stem+'.docx',paragraphs=len(d.paragraphs),tables=len(d.tables),images=len(d.inline_shapes))
if __name__=='__main__':
    report=[build('JEBO_manuscript_v3'),build('JEBO_supplement_v3')];(O/'docx_structure.json').write_text(json.dumps(report,indent=2));print(report)
