"""Extract text from the current manuscript for reproduction-target auditing only."""
from pathlib import Path
from docx import Document

SRC = Path(r"G:\桌面\科研\项目-消费调查\NHB_full_manuscript_draft_2026-09-22_v3_context_abstract.docx")
OUT = Path(__file__).resolve().parent / "manuscript_text_extract.txt"
doc = Document(SRC)
lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
for i, table in enumerate(doc.tables, 1):
    lines.append(f"[TABLE {i}]")
    for row in table.rows:
        lines.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} lines={len(lines)}")
