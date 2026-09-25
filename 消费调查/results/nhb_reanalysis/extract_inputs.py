"""Extract non-confidential questionnaire and Stata metadata for NHB re-analysis."""
from pathlib import Path
import json
import pandas as pd
from docx import Document

ROOT = Path(r"G:\桌面\科研\项目-消费调查")
DOCX = ROOT / "社会心态小调查问卷（整合文字版）(1).docx"
DTA = ROOT / "社会心态小调研数据(1).dta"
OUT = Path(__file__).resolve().parent

doc = Document(DOCX)
lines = []
for p in doc.paragraphs:
    if p.text.strip():
        lines.append(p.text.strip())
for ti, table in enumerate(doc.tables, 1):
    lines.append(f"[TABLE {ti}]")
    for row in table.rows:
        lines.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
(OUT / "questionnaire_text_extract.txt").write_text("\n".join(lines), encoding="utf-8")

reader = pd.io.stata.StataReader(DTA, convert_categoricals=False)
meta = {
    "variable_labels": reader.variable_labels(),
    "value_labels": {
        str(name): {str(key): value for key, value in labels.items()}
        for name, labels in reader.value_labels().items()
    },
}
(OUT / "stata_metadata.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps({"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "variables": len(meta["variable_labels"])}, ensure_ascii=False))
