"""Create manuscript and supplement from saved aggregate results only."""
from pathlib import Path
import json,re
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
O=Path(__file__).resolve().parent;M=O.parent.parent/'manuscript'/'nc_v5';F=M/'figures'
def read(n):return pd.read_csv(O/n)
def pval(x):return '<0.001' if x<.001 else f'{x:.4f}' if .045<x<.055 else f'{x:.3f}'
def ci(r,k=100):return f'{r.estimate*k:.2f} ({r.lo*k:.2f}, {r.hi*k:.2f})'
def newdoc(title):
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85);sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    for e in list(d.styles.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    for e in list(d.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    for s in d.styles:
        if s.type==1:
            s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    n=d.styles['Normal'];n.font.size=Pt(12);n.paragraph_format.line_spacing=2;n.paragraph_format.space_after=Pt(3)
    for name in ['Heading 1','Heading 2','Heading 3']:
        s=d.styles[name];s.font.size=Pt(12);s.font.bold=True;s.font.italic=name!='Heading 1';s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(4);s.paragraph_format.keep_with_next=True;s.paragraph_format.line_spacing=1.15
    d.styles['Title'].font.size=Pt(19);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.line_spacing=1.15;d.styles['Title'].paragraph_format.space_after=Pt(12)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');rr._r.addnext(fld)
    d.core_properties.title=title;d.core_properties.author='Zelin Liu and Yinghao Pan';d.core_properties.subject='Author review draft with unresolved submission information'
    d.add_paragraph(title,'Title');return d
def para(d,t,small=False):
    p=d.add_paragraph(t)
    if small:
        p.paragraph_format.line_spacing=1.15
        for r in p.runs:r.font.size=Pt(10)
    if '[AUTHOR CONFIRMATION' in t:
        from docx.enum.text import WD_COLOR_INDEX
        for r in p.runs:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    return p
def table(d,title,headers,rows,widths=None,note=''):
    d.add_paragraph(title,'Heading 2');t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [6.8/len(headers)]*len(headers)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    for j,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if j==0:
            head=OxmlElement('w:tblHeader');trpr.append(head)
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            tcpr.append(b);mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            if j==0:
                s=OxmlElement('w:shd');s.set(qn('w:fill'),'E8EEF2');tcpr.append(s)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.keep_with_next=(j<len(t.rows)-1 or bool(note))
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9);r.font.bold=j==0
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    if note:para(d,note,True)
    return t
captions={
1:('fig1_distributions','Figure 1. Full response distributions and midpoint-coded averages. a, Unconditional six-category distributions for adults in all nine randomized cells. Shading runs from light for the lowest category to dark for the highest. The first category means essentially no increase. b, Midpoint-coded stated spending shares with pointwise 95% normal intervals based on cell sample variances. Cash, Food and Medical adult cell sizes at RMB 200/1,000/5,000 are 618/592/580, 613/619/599 and 614/611/634. Lines connect independent randomized groups, not within-person trajectories. Source: distribution.csv and profiles.csv.'),
2:('fig2_tail','Figure 2. High-spending probabilities and form contrasts. a, Probability of choosing the >75% response bin in each adult cell; error bars are pointwise 95% normal intervals. Cell sizes are as in Fig. 1. b, Food–Cash and Medical–Cash differences at each amount from saturated cell means with HC3 uncertainty. Thick intervals are pointwise 95%; thin capped intervals are Bonferroni simultaneous 95% intervals for the six displayed contrasts. The Food point estimate crosses Cash but its difference at RMB 5,000 remains uncertain. Source: profiles.csv and form_contrasts.csv.'),
3:('fig3_sensitivity','Figure 3. Sample sensitivity and food-expenditure classification. a, Food–Cash and Medical–Cash slope differences per fivefold amount increase across the full, adult, constant-response-excluded, Q1 and Q2 samples. The samples overlap; the figure is not a sequence of independent replications. b, Food–Cash slope differences among Cash/Food adults whose six-month lower-bound food expenditure exceeds RMB 5,000 and the remaining adults. Classification assumes stable expenditure and comparable eligible purchases. Not confirmed does not mean binding. All bars are pointwise 95% HC3 intervals. The direct between-group triple-interaction test has nominal P=0.0084 and Holm-7 P=0.0504. Source: figure3a_source.csv and mechanism_details.csv.')}
def fig(d,num):
    name,cap=captions[num];d.add_page_break();d.add_picture(str(F/(name+'.png')),width=Inches(6.8));para(d,cap,True)
def main():
    text=(M/'manuscript.md').read_text(encoding='utf8');refs=json.loads((M/'references_original_order.json').read_text(encoding='utf8'));mapping={}
    for m in re.finditer(r'\{cite:([\d,]+)\}',text):
        for s in m.group(1).split(','):
            if s not in mapping:mapping[s]=len(mapping)+1
    (M/'reference_number_map.json').write_text(json.dumps(mapping,indent=2))
    title=text.splitlines()[0][2:];d=newdoc(title)
    for block in text.split('\n\n')[1:]:
        block=block.strip()
        if not block:continue
        if block=='{references}':
            for old,num in mapping.items():para(d,f'{num}. '+refs[old],True)
        elif block=='{table1}':
            d.add_page_break();t=read('main_effects.csv');t=t[t['sample']=='A'];rows=[]
            for y in ['midpoint','top75']:
                for test in ['food-cash','medical-cash','1000-200','5000-200']:
                    r=t[(t.outcome==y)&(t.test==test)].iloc[0];rows.append([('Midpoint' if y=='midpoint' else '>75%')+'\n'+test,ci(r),pval(r.p),pval(r.holm6) if y=='top75' else '—'])
            for test in ['form_omnibus','amount_omnibus']:
                r=t[(t.outcome=='top75')&(t.test==test)].iloc[0];rows.append(['>75% '+test.replace('_',' '),'2 df',pval(r.p),pval(r.holm6)])
            table(d,'Table 1  Marginal effects of the randomized factors',['Outcome and contrast','Estimate and 95% CI (pp)','Nominal P','Holm P'],rows,[2.2,2.3,1.1,1.2],note='N=5,480. Equal-weight contrasts over the other randomized factor; HC3 pointwise intervals. Midpoint values are descriptive approximations to stated spending shares, expressed in percentage points. Holm adjustment applies to the six >75% tests only; midpoint tests are descriptive. Full seven-outcome results are in main_effects.csv. No local family supplies whole-paper error control.')
        elif re.fullmatch(r'\{figure[123]\}',block):fig(d,int(block[-2]))
        elif block.startswith('## '):
            h=block[3:]
            if h in ['References','Methods']:d.add_page_break()
            if h!='Tables and figures':d.add_paragraph(h,'Heading 1')
        elif block.startswith('### '):d.add_paragraph(block[4:],'Heading 2')
        else:
            p=d.add_paragraph();parts=re.split(r'(\{cite:[\d,]+\})',block)
            for part in parts:
                if part.startswith('{cite:'):
                    r=p.add_run(','.join(str(v) for v in sorted(mapping[s] for s in part[6:-1].split(','))));r.font.superscript=True
                else:r=p.add_run(part)
                if '[AUTHOR CONFIRMATION' in part:
                    from docx.enum.text import WD_COLOR_INDEX
                    r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    d.save(M/'NC_MPC_manuscript_v5.docx')
    resolved=re.sub(r'\{cite:([\d,]+)\}',lambda m:'['+','.join(str(mapping[s]) for s in m.group(1).split(','))+']',text)
    resolved=resolved.replace('{references}','\n\n'.join(str(n)+'. '+refs[o] for o,n in mapping.items()))
    (M/'manuscript_resolved.md').write_text(resolved,encoding='utf8')
def supplement():
    d=newdoc('Supplementary information for stated spending responses to transfers across forms and amounts')
    d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.25;d.styles['Normal'].paragraph_format.space_after=Pt(6)
    para(d,'This supplement distinguishes the current fixed analysis from earlier exploratory results on the same respondents. Adults are primary. Historical full-sample results are labelled explicitly. Source tables contain aggregate estimates only; no respondent identifiers or individual response records are released.')
    d.add_paragraph('Supplementary Note 1  Analysis status and theoretical scope','Heading 1')
    para(d,'The study was not preregistered. Earlier analyses examined stated spending levels, fungibility, observable and machine-learning heterogeneity, cross-form prediction, latent traits, amount curves, distributional thresholds, pooled restricted forms, relative scales and approximately 150 moderators. These are analyses of one dataset, not replications. The present bounded plan was fixed after those results were known and before its new estimates were run. A focused family cannot recover an unknown number of prior research choices. Holm-21 is retained even where min-P gives a smaller adjusted P.')
    table(d,'Supplementary Table 1  Competing accounts and discriminating measurements',['Account','Relevant implication','What this study cannot identify'],[
    ['Fungibility under inframarginality','Equal resources can be equivalent if restrictions do not bind and other attributes coincide.','Forms here also differ in expiry, liquidity and explanatory wording.'],
    ['Mechanical restrictions','Effects may depend on transfer size relative to eligible purchases.','Spending bands are proxies; the unclassified group is not demonstrably binding.'],
    ['Labelling and mental accounting','Labels may affect allocation despite feasible substitution.','Labels are not randomized independently of form rules; no direct mediator.'],
    ['Amount-dependent planning or categorization','Size may change planning, perceived spendability or intended allocation.','None of these processes was measured or manipulated. Relative-income fit is not a unique test.']],[1.35,2.55,2.9])
    d.add_page_break();d.add_paragraph('Supplementary Note 2  Outcomes and randomized cells','Heading 1')
    para(d,'The seven codings are ordinal 1–6, midpoint 0/.05/.175/.375/.625/.875, and indicators for categories at or above any spending, 10%, 25%, 50%, and the highest >75% bin. Responses are bounded bins with yuan examples; the highest example is 75–100%, although its verbal label is >75%. Complete adult and full cell distributions and 95% pointwise intervals are supplied in distribution.csv and profiles.csv. Marginal contrasts for every coding are in main_effects.csv. These numerical sources preserve results not emphasized in the narrative.')
    pr=read('profiles.csv');pr=pr[(pr['sample']=='A')&(pr.outcome=='top75')];rows=[]
    dist=read('distribution.csv');dist=dist[dist['sample']=='A']
    for r in pr.itertuples():
        shares=dist[(dist.form==r.form)&(dist.amount==r.amount)].sort_values('category').share.to_numpy()*100
        rows.append([r.form.title(),f'{r.amount:,}',r.N]+[f'{x:.1f}' for x in shares])
    table(d,'Supplementary Table 2  Adult cell distributions',['Form','RMB','N','No rise','<10','10–25','25–50','50–75','>75'],rows,[1.0,.65,.55]+[.7667]*6,note='Shares are percentages. Rounding may prevent a row summing to exactly 100. The first bin means essentially no additional spending.')
    ratio=read('adult_endpoint_ratios.csv');rows=[]
    for f in ['cash','food','medical']:
        r=ratio[(ratio.form==f)&(ratio.scale=='risk_ratio')].iloc[0];o=ratio[(ratio.form==f)&(ratio.scale=='odds_ratio')].iloc[0];rows.append([f.title(),f'{r.estimate:.3f} ({r.lo:.3f}, {r.hi:.3f})',f'{o.estimate:.3f} ({o.lo:.3f}, {o.hi:.3f})'])
    table(d,'Supplementary Table 3  Adult endpoint relative changes',['Form','Risk ratio and 95% CI','Odds ratio and 95% CI'],rows,[1.2,2.8,2.8],note='RMB 5,000 versus RMB 200. Pointwise delta-method intervals on the log scale. These are descriptive within-form comparisons, not multiplicity-adjusted tests that ratios differ across forms.')
    d.add_page_break();d.add_paragraph('Supplementary Note 3  Joint inference','Heading 1')
    para(d,'For each coding, OLS uses an intercept, two form indicators, z and two form×z interactions. HC3 residual influence contributions are stacked across outcomes, giving a 14×14 joint covariance for the interaction coefficients. With covariance V, each simulated centered coefficient vector is drawn as V^(1/2)g, g standard multivariate normal. For each outcome, a 2-df quadratic form supplies the omnibus P and two standardized coefficients supply two-sided scalar P values. The adjusted P is (1 + count of simulated minimum P at most the observed P)/(B+1), B=10,000. Covariance eigenvalues below zero due to numerical precision are truncated at zero. The same covariance estimates both observed and simulated statistics. This centered-error construction targets an asymptotic joint law and permits true and false restrictions together; it is not a permutation under a sharp treatment null.')
    para(d,'HC3 covariances and estimates from the grouped sufficient-statistic implementation were checked against individual-row statsmodels regressions for all seven outcomes. Maximum absolute discrepancies were below 6×10^−16 for coefficients and 4×10^−17 for covariance entries. The historical full-sample focal nominal P values were reproduced. The joint simulation uses unrestricted covariance; it is distinct from the earlier constrained-score diagnostic.')
    t=read('scientific_family21.csv');t=t[t['sample']=='A']
    table(d,'Supplementary Table 4  Adult interaction family',['Outcome','Test','Nominal P','Holm 21','Min P 21'],[[r.outcome,r.test,pval(r.p),pval(r.holm21),pval(r.minP21)] for r in t.itertuples()],[1.55,1.55,1.15,1.25,1.3],note='N=5,480; 21 tests. Outcome selection was post hoc. Complete scalar estimates, pointwise and Bonferroni intervals, plus identical full-sample analyses, are in scientific_family21.csv. Monte Carlo uncertainty at adjusted P≈0.039 is approximately 0.002 standard error with 10,000 draws; a threshold crossing should not be overinterpreted.')
    d.add_page_break();sim=read('simulation_results.csv')
    table(d,'Supplementary Table 5  Complete refit calibration checks',['Scenario','Method','False rejection rate (95% CI)','False null power'],[[r.scenario.replace('_',' '),r.method,f'{r.FWER:.3f} ({r.lo:.3f}, {r.hi:.3f})','—' if pd.isna(r.false_null_mean_power) else f'{r.false_null_mean_power:.3f}'] for r in sim.itertuples()],[2.35,.8,2.5,1.15],note='1,000 independently simulated datasets per scenario, with complete coefficient and HC3 refits and 999 calibration draws each; no failed refits. The partial-null scenario contains 15 true and 6 false restrictions. False-null power is the mean rejection probability across its six false restrictions, not power for the observed focal effect. Binomial intervals are exact. Results do not establish universal finite-sample control or strong power.')
    para(d,'Cell sizes equal the adult observed counts. The pooled null uses overall six-bin frequencies. Additive main effects shift bin probabilities by fixed zero-sum form and dose vectors. The partial-null scenario additionally transfers 0.025×z probability between the third and fourth bins within Food, leaving the highest-bin interaction null. All probabilities and these vectors are explicit in analyze.py. The choice of three scenarios bounds this diagnostic; no scenario or method was selected after seeing simulated rejection rates.')
    para(d,'The historical 1,500-test diagnostic combined 300 models across outcomes, samples, amount parameterizations and links. The same 5,000 saved Gaussian null-score draws reproduce its original max-norm adjustments exactly. Converting each statistic with its own chi-square degrees of freedom gives adjusted P=0.370 for the focal full-sample omnibus, 0.537 for Cash–Food and 0.180 for Cash–Medical (original values 0.414, 0.881 and 0.510). This comparison changes scale calibration, not family content or all null-model assumptions. In the normalized minimum-P draws, df=1, 2 and 4 rows supply the minimum 2,914, 2,032 and 54 times. This is not a measurement of which degrees of freedom dominated the old unnormalized maximum.')
    para(d,'For the original unnormalized maximum, df=4 rows supply 4,949 of 5,000 extremes (98.98%); df=2 rows supply the other 51. Thus the concern about high-df domination is empirically correct for this implementation. The dominant model types are OLS/LPM (2,179 draws) and logit (1,919), rather than the ordered location-scale model (198). The df=4 tests primarily arise from saturated amount interactions. Different links encode different additive nulls. The old ordered location-scale fit zeroed both location and scale interactions under its constrained fit while some reported tests involved location only. Consequently the historical diagnostic is not presented as a uniformly calibrated strong-FWER procedure. The prior 7/200 check concerned a fixed-score family maximum under a pooled null, not one isolated test; it nevertheless did not refit the full estimation procedure or establish power.')
    d.add_page_break();d.add_paragraph('Supplementary Note 4  Direct diagnostic tests','Heading 1')
    mech=read('mechanism_tests.csv')
    table(d,'Supplementary Table 6  Seven diagnostic restrictions',['Restriction','N','df','Nominal P','Holm 7'],[[r.test.replace('_',' '),r.N,r.df,pval(r.p),pval(r.holm7)] for r in mech.itertuples()],[3.1,.75,.45,1.25,1.25],note='Adult >75% outcome; M1 income mapping. Complete coefficient estimates and pointwise intervals where scalar are supplied in mechanism_tests.csv; subgroup contrasts and alternative mapping results are in mechanism_details.csv.')
    para(d,'For nominal versus relative income scaling, the unrestricted model is logit P(Y=1)=form-specific intercept + form-specific slope×ln(amount) + form-specific slope×ln(income). ABS sets the three income slopes to zero. REL sets each income slope plus its amount slope to zero. These are separately testable restrictions of the same unrestricted model; ABS and REL themselves are nonnested. Income triple interactions add all lower-order terms before jointly testing the two three-way coefficients. M2 gives nominal restriction P=0.0059, ratio restriction P=0.0375 and income triple-interaction P=0.444 before correction; it is a descriptive mapping sensitivity.')
    para(d,'The food classification model includes all lower-order terms in form×z×G, restricted to Cash and Food. G=1 when six times the reported food-expenditure lower bound exceeds 5,000. The ratio decomposition includes form-specific ln(amount) and ln(food lower bound) slopes, then tests whether their Food–Cash differences sum to zero. It uses 3,476 adults with positive lower bounds; G analysis uses all 3,621 Cash/Food adults. The ratio is not a treatment and the expenditure bands need not capture eligible purchases over the future validity period.')
    det=read('mechanism_details.csv');ss=det[det.test.str.startswith('Q')]
    table(d,'Supplementary Table 7  Passing and failing response style screens',['Screen and group contrast','Group N','Slope difference and 95% CI (pp)'],[[r.test.replace(' 0 ',' fail ').replace(' 1 ',' pass '),int(r.subgroup_N),ci(r)] for r in ss.itertuples()],[3.0,1.0,2.8],note='Contrasts are Food–Cash or Medical–Cash per fivefold increase. All groups are subsets of adults. Pointwise HC3 intervals; group-specific P values are not interpreted as evidence of between-group differences. Direct joint differences are tested in Table 6. Nine-cell group counts and events are in screen_cell_counts.csv.')
    d.add_page_break();d.add_paragraph('Supplementary Note 5  Prediction and income coding','Heading 1')
    para(d,'M1 maps current income codes 1–6 to 250, 750, 2,000.5, 5,500.5, 11,500.5 and 20,000 RMB. Legacy codes 11–16 map to 500, 2,000.5, 4,000.5, 6,500.5, 10,000.5 and 16,000.5 RMB. M2 changes only codes 1, 6 and 11 to 125, 30,000 and 250. These are band representations rather than exact incomes. Fold allocation is reconstructed on the full sample using five cell-stratified folds and seed 20261003, then restricted to adults. All five candidate models share these folds. Logistic optimizers converged in every fold; only training data determine column scaling.')
    pred=read('prediction_scores.csv')
    table(d,'Supplementary Table 8  Held out prediction',['Mapping','Model','Log loss','Gain vs ABS (%)','Gain vs additive (%)'],[[r.mapping,r.model,f'{r.logloss:.6f}',f'{100*r.relative_gain_vs_ABS:.3f}',f'{100*r.relative_gain_vs_additive:.3f}'] for r in pred.itertuples()],[.75,1.8,1.3,1.5,1.45],note='Positive gain means smaller loss. Fold-level losses and sample sizes are in prediction_folds.csv. Differences are descriptive; the common training sets preclude treating individual held-out losses as independent for an ordinary paired inferential interval. No materiality threshold selects a mechanism.')
    para(d,'Historical Cash descriptive associations use the full Cash sample (n=1,798) and jointly adjust for dose, harmonized income rank and emergency fundraising. They were previously labelled positive controls; that label is withdrawn. They cannot independently validate actual spending. Historical moderator and 150-variable analyses are retained in source files for disclosure, without renewed selection. Their null findings are limited by precision: for example, the income-rank moderation MDE for Cash–Food on the high-spending outcome is about 3.34 points per SD, larger than the observed average slope contrast. This is a plug-in normal-design MDE, not a simulated power guarantee.')
    d.add_page_break();d.add_paragraph('Supplementary Note 6  Census calibration','Heading 1')
    para(d,'The official source is National Bureau of Statistics Table A0401, https://www.stats.gov.cn/sj/pcsj/rkpc/7rp/zk/html/A0401.jpg, accessed 3 October 2026. A visual transcription of ages 18, 19, 20–24, 25–29 and subsequent five-year bands supplies census_2020_transcription.csv. Tertiary counts sum junior college, bachelor, masters and doctorate. Lower education is total minus high school and tertiary. Target proportions concern the 2020 population, not the current population or survey recruitment frame. Adult sample education codes 1/2/3–5 map to lower/high/tertiary. Coarse cells were chosen without outcome values. No claim is made that the population above age 81 is represented.')
    cc=read('calibration_cells.csv')
    table(d,'Supplementary Table 9  Age by education calibration',['Age','Education','Sample N','Sample %','Target %','Capped %'],[[r.agegroup,r.education,r.sample_N,f'{100*r.sample_share:.2f}',f'{100*r.target_share:.2f}',f'{100*r.cap10_share:.2f}'] for r in cc.itertuples()],[.85,1.15,.8,1.25,1.25,1.5])
    cd=read('calibration_diagnostics.csv')
    table(d,'Supplementary Table 10  Weight diagnostics',['Scheme','ESS','Maximum weight','Largest cell gap (pp)'],[[r.scheme,f'{r.ESS:.1f}',f'{r.max_weight:.2f}',f'{r.max_abs_target_discrepancy*100:.2f}'] for r in cd.itertuples()],[1.65,1.05,1.9,2.2],note='Weights have mean one. Capped weights obey a final normalized maximum of 10; 2.92% of respondents attain that cap. Exact matching and this bound cannot both be achieved with the observed composition. Conditional-weight HC3 estimates, pointwise intervals and local Holm-3 adjustments are in calibration_effects.csv. These diagnostic P values do not replace the primary Holm-21 inference.')
    d.add_page_break();d.add_paragraph('Supplementary Note 7  Specification sensitivity','Heading 1')
    para(d,'The next plot puts linear probability estimates on a common endpoint scale. A trend slope difference is multiplied by two; a saturated model uses its RMB 5,000-versus-200 interaction; an endpoint model excludes the middle amount. The saturated and endpoint point estimates coincide by construction, with potentially different HC3 estimates from leverage. Across five overlapping samples these are 15 displayed specifications per contrast, not 15 independent datasets. No significance-count or pooled median test is used, because the samples also change the target population.')
    d.add_picture(str(F/'figS1_specifications.png'),width=Inches(6.8));para(d,'Supplementary Figure 1. Common-scale specification display. Points are Food–Cash or Medical–Cash differences in the RMB 200-to-5,000 probability change, with pointwise 95% HC3 intervals. R, full; A, adults; C, constant-response and duplicate-ID exclusions; Q1/Q2, response-style screens defined in Methods. Source: spec_curve_source.csv.',True)
    d.add_page_break();d.add_picture(str(F/'figS2_calibration.png'),width=Inches(6.8));para(d,'Supplementary Figure 2. Age-by-education calibration sensitivity. Points show Food–Cash and Medical–Cash trend differences per fivefold amount increase. Bars are pointwise 95% HC3 intervals conditional on the displayed weights. ESS denotes the Kish effective sample size; it does not measure correction of selection bias. Source: calibration_effects.csv.',True)
    para(d,'All local families answer distinct questions and do not guarantee whole-paper familywise control. Frozen definitions, simulation settings and input hashes are in freeze.json. Historical models, selection screens, cross-threshold analyses and analysis history remain available alongside the current results. No new moderator, outcome threshold, income mapping, Bayesian prior search or independent-data claim was added after the current plan. Unresolved recruitment and ethics information is tracked separately for the authors, and this package is not certified submission-ready.')
    # Continuous supplementary flow avoids nearly empty continuation pages.
    drawing_index=0
    for pp in list(d.paragraphs):
        breaks=pp._p.xpath('.//w:br[@w:type="page"]')
        if breaks and not pp.text.strip():pp._p.getparent().remove(pp._p)
        if pp.text.startswith('Supplementary Note 7'):pp.paragraph_format.page_break_before=True
        if pp._p.xpath('.//w:drawing'):
            pp.paragraph_format.keep_with_next=True;pp.paragraph_format.page_break_before=drawing_index>0;drawing_index+=1
    d.save(M/'NC_MPC_supplement_v5.docx')
main();supplement();print('Created manuscript and supplementary Word documents')
