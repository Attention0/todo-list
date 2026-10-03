"""Render the two resolved Markdown manuscripts into native Word structures."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
O=Path(__file__).resolve().parent;M=O.parents[1]/'manuscript'/'jebo_v1'
def tidy(t):
 if t.startswith('http'):return t
 t=t.replace('TableS','Table S').replace('Figure1','Figure 1').replace('Table3','Table 3')
 # Editorial spacing in compact audit prose; preserve URLs and file names.
 bits=re.split(r'(https?://\S+|[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}|\b\S+\.(?:csv|md|py|json)\b)',t)
 for i in range(0,len(bits),2):
  s=bits[i];s=re.sub(r'(?<=[a-z])(?=\d)',' ',s);s=re.sub(r'(?<=\d)(?=[A-Za-z])',' ',s);s=re.sub(r'(?<=%)(?=[A-Za-z])',' ',s);s=re.sub(r'(?<=\d)(?=and\b)',' ',s);bits[i]=s
 return ''.join(bits)
def runtext(p,t):
 # Plain human text; Markdown emphasis is handled without leaking markers.
 for j,x in enumerate(re.split(r'\*\*(.*?)\*\*',t)):
  r=p.add_run(x);r.bold=bool(j%2)
def newdoc():
 d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=s.bottom_margin=Inches(.78);s.left_margin=s.right_margin=Inches(.82)
 for style in d.styles:
  if style.type==1:style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
 normal=d.styles['Normal'];normal.font.size=Pt(11.5);normal.paragraph_format.line_spacing=1.23;normal.paragraph_format.space_after=Pt(7)
 for name,size in [('Title',19),('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
  st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.line_spacing=1.1;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
 d.styles['Title'].paragraph_format.space_before=Pt(0)
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 d.core_properties.author='Zelin Liu and Yinghao Pan';d.core_properties.subject='Spending responses to hypothetical transfer scenarios'
 return d
def make_table(d,lines):
 rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|[\s:|\-]+\|$',line)]
 n=len(rows[0]);t=d.add_table(rows=0,cols=n);t.autofit=False
 # More width for meaningful labels, compact numeric columns.
 weights=([1.0,.65,.45]+[.7]*6) if n==9 else ([1.25,.6]+[.8]*(n-2)) if n>=7 else ([1.2,1.6,2.4,1,1] if n==5 else [1]+[1.65]*(n-1))
 if rows[0][0]=='Form and use rights':weights=[3.4,1,1,1]
 if rows[0][0]=='Form' and n==5:weights=[.8,.7,1.1,1.1,2.3]
 if rows[0][0].startswith('Screen /'):weights=[2.5,.8,2.7]
 if rows[0][0]=='Restriction':weights=[2.5,.65,.5,1,1]
 widths=[6.63*x/sum(weights) for x in weights]
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for j,row in enumerate(rows):
  cells=t.add_row().cells;pr=cells[0]._tc.getparent().get_or_add_trPr();cs=OxmlElement('w:cantSplit');pr.append(cs)
  if j==0:pr.append(OxmlElement('w:tblHeader'))
  for i,(c,v) in enumerate(zip(cells,row)):
   c.width=Inches(widths[i]);c.text=tidy(v);tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');mar.append(el)
   tcpr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','bottom']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D5D5D5');borders.append(el)
   tcpr.append(borders)
   if j==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EFEFEF');tcpr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    # Keep short tables with their notes; long tables repeat headers and retain
    # at least two opening data rows instead of stranding a one-row fragment.
    p.paragraph_format.keep_with_next=len(rows)<=13 or j<3 or j==len(rows)-1
    for r in p.runs:r.font.size=Pt(9 if n<7 else 8);r.bold=j==0
 return t
def build(stem):
 src=M/(stem+'.md');text=src.read_text(encoding='utf8');text='\n\n'.join(tidy(b) if not b.startswith('![') else b for b in text.split('\n\n')).rstrip()+'\n';src.write_text(text,encoding='utf8')
 d=newdoc();blocks=text.split('\n\n');references=False;intables=False
 for b in blocks:
  b=b.strip()
  if not b:continue
  if b.startswith('# '):
   title=b[2:];d.add_paragraph(title,'Title');d.core_properties.title=title
  elif b.startswith('## '):
   name=b[3:];references=name=='References';intables=name=='Tables and figures'
   if intables:continue
   p=d.add_paragraph(name,'Heading 1')
   if references or intables:p.paragraph_format.page_break_before=True
  elif b.startswith('### '):
   p=d.add_paragraph(b[4:],'Heading 2')
   if intables:p.paragraph_format.page_break_before=True
  elif b.startswith('|'):make_table(d,b.splitlines())
  elif b.startswith('!['):
   match=re.match(r'!\[(.*?)\]\((.*?)\)',b);p=d.add_paragraph();p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
   rr=p.add_run();rr.add_picture(str(M/match.group(2)),width=Inches(6.6));docpr=rr._r.xpath('.//wp:docPr')[0];docpr.set('descr',match.group(1))
  else:
   p=d.add_paragraph();runtext(p,b)
   if references or re.match(r'^Figure \d+\.',b):
    p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(10)
   if references:p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
 # Remove decorative paragraph borders inherited from any default style.
 for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.save(M/(stem+'.docx'))
 return dict(file=stem+'.docx',paragraphs=len(d.paragraphs),tables=len(d.tables),images=len(d.inline_shapes))
if __name__=='__main__':
 r=[build('JEBO_manuscript_v1'),build('JEBO_supplement_v1')];(O/'docx_structure.json').write_text(json.dumps(r,indent=2));print(r)
