"""Independent numerical and document checks; aggregate data only."""
from pathlib import Path
import csv,json,hashlib,re,math,subprocess
import numpy as np
from docx import Document
from pypdf import PdfReader
from build_docx import tidy
O=Path(__file__).resolve().parent;P=O.parent/'nc_evidence_revision';ROOT=O.parents[1];M=ROOT/'manuscript'/'jebo_v1'
def read(p):
 with Path(p).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
checks={};old=json.loads((O/'historical_input_hashes.json').read_text(encoding='utf8'))
changed=[n for n,h in old.items() if sha(ROOT/n)!=h];assert not changed,changed;checks['historical_files_unchanged']=len(old)
c=read(O/'extensive_intensive_cells.csv');assert len(c)==9 and sum(int(r['N']) for r in c)==5480
src=[r for r in read(P/'distribution.csv') if r['sample']=='A'];prof=[r for r in read(P/'profiles.csv') if r['sample']=='A']
scores={'ordinal':np.arange(1,7),'midpoint':np.array([0,.05,.175,.375,.625,.875]),'any_spending':np.array([0,1,1,1,1,1]),'ge10':np.array([0,0,1,1,1,1]),'ge25':np.array([0,0,0,1,1,1]),'ge50':np.array([0,0,0,0,1,1]),'top75':np.array([0,0,0,0,0,1])}
counts={};maxerr=0
for r in c:
 key=(r['form'],int(r['amount']));cnt=np.array([int(r['count'+str(j)]) for j in range(1,7)]);counts[key]=cnt
 oldcnt=np.array([int(float(q['count'])) for q in sorted([q for q in src if q['form']==key[0] and int(q['amount'])==key[1]],key=lambda q:int(q['category']))]);assert np.array_equal(cnt,oldcnt)
 assert abs(float(r['midpoint'])-float(r['p_any'])*float(r['conditional_midpoint']))<1e-14
 for y,sc in scores.items():
  mean=cnt@sc/cnt.sum();q=next(q for q in prof if q['form']==key[0] and int(q['amount'])==key[1] and q['outcome']==y);maxerr=max(maxerr,abs(mean-float(q['estimate'])))
assert maxerr<1e-13;checks['all_63_profile_means_max_error']=maxerr
# Independent grouped individual-observation HC3 slope calculation.
fit={};errs=[]
for y,sc in scores.items():
 for f in ['cash','food','medical']:
  ns=np.array([counts[(f,a)].sum() for a in [200,1000,5000]]);X=np.array([[1,0],[1,1],[1,2]],float);inv=np.linalg.inv(X.T@(ns[:,None]*X));sy=np.array([counts[(f,a)]@sc for a in [200,1000,5000]]);beta=inv@X.T@sy
  h=np.einsum('ij,jk,ik->i',X,inv,X);rss=np.array([counts[(f,a)]@((sc-X[i]@beta)**2) for i,a in enumerate([200,1000,5000])]);V=inv@(X.T@((rss/(1-h)**2)[:,None]*X))@inv
  fit[(y,f)]=(beta[1],V[1,1])
  q=next(q for q in read(P/'within_form_slopes.csv') if q['sample']=='A' and q['outcome']==y and q['test']==f);errs += [abs(beta[1]-float(q['estimate'])),abs(math.sqrt(V[1,1])-float(q['se']))]
assert max(errs)<1e-12;checks['independent_21_within_form_HC3_max_error']=max(errs)
fam=[r for r in read(P/'scientific_family21.csv') if r['sample']=='A'];ps=[]
for r in fam:
 y=r['outcome'];bc,vc=fit[(y,'cash')];bf,vf=fit[(y,'food')];bm,vm=fit[(y,'medical')]
 if r['test']=='omnibus':v=np.array([[vc+vf,vc],[vc,vc+vm]]);b=np.array([bc-bf,bc-bm]);p=math.exp(-float(b@np.linalg.solve(v,b))/2)
 else:
  other=r['test'].split('-')[1];b,v=fit[(y,other)];est=bc-b;se=math.sqrt(vc+v);assert abs(est-float(r['estimate']))<1e-12 and abs(se-float(r['se']))<1e-12;p=math.erfc(abs(est/se)/math.sqrt(2))
 assert abs(p-float(r['p']))<1e-11;ps.append(p)
order=np.argsort(ps);adj=np.empty(21);cum=0
for rank,i in enumerate(order):cum=max(cum,(21-rank)*ps[i]);adj[i]=min(1,cum)
assert max(abs(adj[i]-float(r['holm21'])) for i,r in enumerate(fam))<1e-11;checks['21_interaction_nominal_and_Holm_independent']=True
for r in read(O/'extensive_intensive_decomposition.csv'):
 lo=next(q for q in c if q['form']==r['form'] and q['amount']=='200');hi=next(q for q in c if q['form']==r['form'] and q['amount']=='5000')
 assert abs(float(r['extensive'])+float(r['intensive'])-float(r['total']))<1e-14
 assert abs(float(r['total'])-(float(hi['midpoint'])-float(lo['midpoint'])))<1e-14
checks['adult_cells_and_decomposition_identities']=True
meta=json.loads((O/'decomposition_verification.json').read_text());assert meta['draws']==4000 and meta['seed']==2026100316 and meta['invalid_zero_positive_draws']==0 and not meta['raw_data_access']
assert meta['source_sha256']==sha(P/'distribution.csv') and meta['manifest_sha256']==sha(O/'JEBO_ANALYSIS_MANIFEST.md');checks['bootstrap_settings_and_manifest_hash']=True
texts={n:(M/(n+'.md')).read_text(encoding='utf8') for n in ['JEBO_manuscript_v1','JEBO_supplement_v1']}
assert not any(re.search(r'\{(?:table|figure|s\d|references)',t) for t in texts.values())
main=texts['JEBO_manuscript_v1'];intro=main.split('## 1. Introduction')[1].split('## 2.')[0];abstract=main.split('## Abstract')[1].split('Keywords:')[0]
assert 150<=len(abstract.split())<=220 and 1500<=len(intro.split())<=2000
checks['abstract_words']=len(abstract.split());checks['intro_words']=len(intro.split())
assert len(read(O/'jebo_literature_matrix.csv'))==54;assert len(read(O/'jebo_outlet_benchmark.csv'))==13
checks['literature_matrix_and_outlet_counts']=[54,13]
assert '1985.0' not in main and '2026.0' not in main
assert not any(s in main.split('## References')[0] for s in ['PR16','PR #','all-X','SHAP','frozen analysis','NC v5','Nature Communications','change of outlet'])
checks['main_internal_process_language_absent']=True
for stem,t in texts.items():
 d=Document(M/(stem+'.docx'));assert len(d.tables)==(3 if 'manuscript' in stem else 15);assert len(d.inline_shapes)==(3 if 'manuscript' in stem else 0)
 assert d.styles['Title'].font.color.rgb==__import__('docx').shared.RGBColor(0,0,0)
 p=M/'qa'/'final'/stem/(stem+'.pdf');pages=PdfReader(p).pages;png=list((M/'qa'/'final'/stem).glob('page-*.png'));assert len(png)==len(pages),(stem,len(png),len(pages))
 pagesummary=[]
 for i,page in enumerate(pages):
  txt=page.extract_text() or '';assert len(txt.strip())>15,(stem,i);pagesummary.append(dict(page=i+1,words=len(txt.split()),start=txt[:100].replace('\n',' ')))
 checks[stem]=dict(pages=len(pages),tables=len(d.tables),images=len(d.inline_shapes),page_summary=pagesummary)
# Make ledger typography identical to the final Markdown, retaining sources/status.
ledger=read(M/'JEBO_CLAIM_LEDGER.csv')
for r in ledger:r['claim_text']=tidy(r['claim_text']).rstrip();assert r['claim_text'] in texts[r['location'].split(' :: ')[0].replace('.md','')],r['location']
with (M/'JEBO_CLAIM_LEDGER.csv').open('w',encoding='utf8',newline='') as f:w=csv.DictWriter(f,fieldnames=list(ledger[0]));w.writeheader();w.writerows(ledger)
checks['claim_blocks_matched_to_final_text']=len(ledger)
# All source data for the figures are adult aggregate records.
for file in (M/'figures').glob('*source.csv'):assert all(r['sample']=='A' for r in read(file))
for file in (M/'figures').glob('*.pdf'):assert len(PdfReader(file).pages)==1
checks['figures_adult_only_vector_single_page']=True
(O/'verification_results.json').write_text(json.dumps(checks,indent=2),encoding='utf8');print(json.dumps({k:v for k,v in checks.items() if not k.startswith('JEBO_')},indent=2))
