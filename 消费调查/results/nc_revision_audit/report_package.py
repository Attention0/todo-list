"""Build audit documents from aggregate outputs; never edit the manuscript."""
from pathlib import Path
import hashlib,json,importlib.metadata,re
import pandas as pd
from docx import Document
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
LOCAL=Path('G:/桌面/科研/项目-消费调查')
RAW=LOCAL/'社会心态小调研数据(1).dta'
QUESTION=LOCAL/'社会心态小调查问卷（整合文字版）(1).docx'
MAN=LOCAL/'NC_MPC_manuscript_draft_v2.docx'
def put(name,text): (OUT/name).write_text(text.strip()+'\n',encoding='utf-8')
def read(name):return pd.read_csv(OUT/name)
def md(t):
    t=t.copy().fillna('—'); cols=list(t.columns)
    def f(v):return f'{v:.6g}' if isinstance(v,float) else str(v).replace('|','/').replace('\n',' ')
    return '| '+' | '.join(cols)+' |\n| '+' | '.join(['---']*len(cols))+' |\n'+'\n'.join('| '+' | '.join(f(v) for v in row)+' |' for row in t.itertuples(index=False,name=None))
def main():
    a=read('table_A1_focal_estimand.csv'); profiles=read('form_amount_profiles.csv')
    q=read('quality_gradient_source.csv'); binding=read('bindingness_ratio_tests.csv')
    cv=read('relative_scale_cv.csv'); mde=read('moderator_power_mde.csv'); pc=read('positive_controls.csv')
    put('NC_REVISION_RESULT.md',f'''# NC revision evidence audit — decision memo

Date: 2026-10-03. All findings concern hypothetical stated additional total spending, not realized expenditure. This revision-stage finite family follows the saved manifest; neither the original study nor the accumulated story was preregistered. No manuscript text was changed.

## Eight required answers

1. **Global correction: NO.** R Top75 three-form equal-slope Wald=9.6480 (2 df), raw P=.00803447, global max-T P=.41371726. Cash−Food=-2.81935 pp per fivefold increase, 95% CI [-5.09262,-0.54608], raw P=.0150645, global P=.88062388. Cash−Medical=-3.08952 pp, CI [-5.12198,-1.05705], raw P=.00288847, global P=.51009798. These are pointwise, not simultaneous, CIs. See full focal table below. No scalar CI exists for the omnibus; its two-component estimate/CI vector is in the CSV.
2. **Separate forms:** Cash Top75 slope=-4.05460 pp; Food=-1.23524 pp; Medical=-.96508 pp per fivefold. Cash has the largest observed upper-tail compression. Midpoint and ordinal form profiles below show the mean pattern separately. Food and Medical cannot be asserted identical: their slope contrast P=.790906 is not equivalence. No fourth randomized Restricted arm exists.
3. **Simple bindingness: incomplete, not supported as a sufficient explanation.** Conditional on amount, Food−Cash log-r Top75 gradient=+.0196842 (CI [.0060452,.0333232]), raw P=.0046733, Holm across 75 C tests=.322456. Thus observed Cash−Food gap declines, rather than increases, with this baseline-dependent ratio. This is an association, not a causal ratio effect. Fixed strictly inframarginal sample n=2815: Food−Cash slope difference=+.0123031 (CI [-.0144757,.0390818], P=.3679); pooled equal-amount level difference=-.0117278 (CI [-.0328810,.0094253], P=.2772). The slope mechanism remains unresolved; a non-significant level is not equivalence.
4. **Relative-scale categorization: unsupported by the focal prediction comparison; direct process evidence NOT FEASIBLE.** R Top75 REL improvements are .1601%, .1114%, .0328% (M1/M2/rank), all below the revision-stage 1% criterion and all paired loss CIs include zero. Secondary ordinal/midpoint improvements are also below 1%. These do not establish income-based collapse. Questionnaire has no intended-use, planning-horizon, spendability or earmarking process response. Scale-dependent categorization remains a Discussion hypothesis only.
5. **Broadly shared tendency: cannot claim.** Four theory moderators × three outcomes × three targets: no Holm-36 rejection. Top75 Cash-slope 80% MDE is about 2.48–2.93 pp/1SD; Cash−Food differential MDE 3.34–3.88 pp exceeds the average 2.82 pp contrast; Cash−Medical MDE 3.08–3.59 pp is comparable to/exceeds its 3.09 pp contrast. The design cannot rule out moderation on the scale of the average interaction. Normal-design MDE uses actual HC3 uncertainty, not a simulated exact power guarantee.
6. **Positive controls: MIXED.** Cash size gradient is negative and survives Holm-9 for Top75. Higher fundraising capacity predicts lower Top75, raw P=.02097 but Holm P=.06291. Income rank predicts *higher* Top75 (estimate +.0240527, CI [.0097183,.0383871], Holm P=.006036), opposite the frozen expected sign. Income is not liquid wealth and the literature direction is not universal. This is limited construct-validity evidence, never validation of stated as realized MPC.
7. **Quality/weighting:** same-signed Cash−Food/Cash−Medical contrasts across R/A/C/Q1/Q2; omnibus raw P=.00803/.00903/.00364/.15392/.41061; global P=.41372/.44411/.25815/.99360/1.00000. Smaller quality samples have wider uncertainty; their composition differs. Weighting NOT FEASIBLE in this run because compatible official adult age×education margins could not be verified/retrieved; unweighted robustness is not a weighting test. No nationally representative claim. See status files.
8. **Recommended strength: B — interesting but should be framed as suggestive.**

## Rationale (eight sentences)

**Fact:** the observed Cash Top75 decline exceeds the Food and Medical declines. **Fact:** the focal omnibus and both primary contrasts fail the finite-family global correction. **Inference:** the form-by-size pattern remains a useful descriptive candidate, not strong corrected evidence. **Fact:** ordered location-scale sensitivity does not restore a robust full-distribution interaction. **Elimination:** the available bindingness and relative-scale tests do not establish either mechanism, but cannot exclude every version of them. **Inference:** null moderation is limited by differential-slope MDEs and cannot support universality. **Fact:** positive controls are mixed and ethics/minor coverage and recruitment details remain unresolved. **Interpretation:** a bounded account of stated response patterns is warranted; endogenous earmarking is speculation and no further exploratory analysis follows this package.

## Focal table — probability units, not percentage points

{md(a[['estimand','estimate','lo','hi','p','global_maxT_p','N','df']])}

## Separate-form profiles — all three outcomes

{md(profiles[profiles.outcome.isin(['top75','midpoint','ordinal'])][['form','amount','outcome','N','estimate','lo','hi']])}

## Interpretation boundaries

- **Fact:** 300/300 fitted models and 1,500/1,500 contrast rows were estimable. Global family covers seven outcomes, five samples, three dose representations and appropriate links. 5,000 centered constrained-null score-multiplier draws; zero draw-level optimizer refits. Null score-level simulation: 7/200 rejections (3.5%; 95% binomial CI 1.42–7.08%). This is an asymptotic joint resampling check, not exact randomization inference or a full model-refit bootstrap.
- **Fact:** ordered-model primary ordinal omnibus raw P spans .078–.118 for the nonlinear models; global P>.94. Location and scale are allowed to vary by form/dose; no strong location effect after correction. The alternative alone cannot prove which scale mechanism generated the ordinal pattern.
- **Fact:** at ¥5000 only Cash−Medical and Restricted−Cash meet Holm-12 TOST at the widest ±5pp margin; Food−Medical does not. No independent SESOI was available. **Inference:** do not claim convergence or equal arms; wide-margin descriptive sensitivity is weaker than substantive equivalence.
- **Fact:** all-X supplemental six primary families have zero BH q<.10 and zero q<.05. Old analysis is not an independent replication. No nominal winner is highlighted.
- **Inference:** upper-tail visibility does not establish tail-specificity. No new cross-threshold claim is introduced here; all threshold probabilities, endpoint RR/OR and probability AMEs are disclosed.
- **SUBMISSION BLOCKER:** current documentation does not establish ethics/consent coverage of 17 minors. Recommend adult-only primary sample n=5480 unless authors supply appropriate coverage; analysis manifest is unchanged and R remains disclosed.
- **NOT FEASIBLE:** response-time/process measures and adult-compatible calibration in this run. Optional Bayesian model was not run. No substitute proxy, new moderator, outcome, tuned threshold or manuscript rewrite.

## Audit navigation

Methods/limits: maxT_diagnostics.md; analysis_history.md; wording_identification_note.md; relative_scale_falsification_memo.md. Mechanisms: mechanism_prediction_table.md; questionnaire_mechanism_audit.md; response_time_audit.md. Submission: ethics_minors_audit.md; author_information_needed.md; open_science_release_checklist.md; NC_submission_requirements.md. Literature: reference_audit.csv; NC_revision_literature_map.md; revised_intro_outline.md. Verification: verification_results.json; figure_audit.md. Reproduction/publication details: README.md; RESULT.md.
''')
    put('analysis_history.md','''# Analysis history

This study was not preregistered. Before this round investigators inspected stated-MPC levels, randomized form contrasts, fungibility, theory and ML/honest HTE, cross-form prediction, latent traits, amount curves, distribution thresholds, Restricted pooling, headroom/relative scales, limited multiverses, who-drives moderators, all-X multi-outcome screens, and main/Hero figure presentations. The current form-by-size story is post hoc. Historical findings reuse this dataset and are not independent replications.

Main was updated to b5f26e5. Existing figure/analysis lineage is inherited from feature/mpc-hero-figure. Revision manifest was saved and committed at a80b94e before new model runs, dated 2026-10-03; SHA256 6b1e95952886d4cc0a245579d03fe72b40e79f7c925669cd1f5ba40c9cc87375. This is a revision-stage freeze, not preregistration. Seven outcomes/five samples/three dose representations/appropriate model families/five contrasts give 300 models and 1,500 tests. Manifest records moderator variables, 1% CV materiality threshold, TOST margins and multiplicity families.

See bug_log.md for export/filter/reference-retrieval fixes; none added an estimand or specification. Income schemes were distinguished using raw labels within the saved mapping plan; historical results were not edited. The optional Bayesian model was explicitly skipped before analysis. No post-result additions. This finite audit ends here.
''')
    put('maxT_diagnostics.md','''# Constrained-null joint max-T diagnostics

Finite family: 300 models, 1,500 test rows, all estimable; no failed observed/null model fits. Convergence gradient and minimum Hessian eigenvalues are in model_diagnostics.csv. OLS/LPM uses individual HC3. Binary/ordered links use individual score–Hessian sandwich HC1. Aggregated likelihood cells are sufficient groups, NOT clusters. Saturated pairwise tests are 2-component vectors; omnibus is 4df; trend/endpoint pairwise are scalar, omnibus 2df. Full parameterization and contrasts are in grid_models.py and the manifest.

For each specification the null retains form and amount main effects and removes interaction terms (location and scale interactions in the heterogeneous-choice model). Centered individual score contributions use null parameters, the fitted null Hessian, and independent Gaussian person multipliers. All outcomes, samples and models share each respondent's multiplier. Compression into cell×sample-membership×ordinal-bin sufficient groups is exact for the multiplier process because scores within these groups coincide. Full-score nuisance components are retained in covariance/influence calculations.

Draw maximum includes |standardized scalar contrast| and sqrt(Wald) Euclidean norms for multi-df contrasts. This explicitly conservative heterogeneous-df norm family is not a scalar-t-only test family. Adjusted P=(1+#max>=observed norm)/5001. Covariance/nuisance parameters stay fixed across draws; **no per-draw optimizer refitting**. Thus no draw-level nonestimable fits exist; this does not mean 5,000×300 fitted optimizations succeeded. 5,000 shared draws, seed 20261003; all maxima exported. This is a suitable asymptotic centered-score multiplier correction, not exact randomization inference or a parametric model-refit bootstrap. Its validity depends on correct null score centering, design independence and asymptotic linearization, and weak links/ordinal cutpoints can add finite-sample error.

Null check: fixed observed cell×quality-membership sizes, independent multinomial six-bin responses using overall observed category probabilities, hence no interactions. A separate 5,000-draw null-score calibration (seed20261004) gives the 95th percentile. 200 independently simulated categorical datasets (seed20261005) are mapped through fixed null-score influence matrices, NOT refitted models. Rejections=7/200=3.5%; exact binomial 95% CI [.0141855,.0707810] contains .05. This supports approximate size for the score-level family under this one pooled null; it does not validate all possible heterogeneous nulls or full optimizer refit size. Counts and all 200 maxima are exported. Omnibus global P=.413717, Cash−Food=.880624, Cash−Medical=.510098. No alternative family was used to rescue these results.
''')
    put('wording_identification_note.md','''# Wording and identification

**Fact:** Cash is unrestricted, paid to bank/WeChat/Alipay; Food is food/daily-goods-only, expires in six months; Medical accumulates long-term and covers self/family medical costs. These bundled attributes differ; one cannot identify pure restriction, label or expiry separately.

All arms ask “您的【总消费】预计会比原计划多花多少钱？” Cash category1 says savings/debt of the subsidy; Food category1 describes buying planned items and saving/repaying freed cash; Medical describes paying planned medical expenses and saving/repaying freed cash. Consequently cross-form no-spending gaps are wording-sensitive. Within each form the scenario and category explanations remain fixed across amounts except the assigned number and arithmetic bin examples. **Inference:** within-form gradients and differences of gradients are less exposed to *fixed additive* wording differences, conditional on no wording×amount effect; this assumption is not directly tested. Within-Cash 200→5000 is more defensible than cross-form category1 comparisons. No rewrite or new outcome. Top75 always original category6; midpoint response coding is [0,.05,.175,.375,.625,.875], not observed continuous MPC. “Most visible in the upper tail” may be descriptive; “tail-specific” is not justified by threshold-wise significance comparisons.
''')
    mech=[
    ['Full fungibility / fully inframarginal','Cash−Food level=0 under equivalent budgets, horizons, use and no label effects','Equal amount slopes under same conditions','Benchmark conditions strong; non-significance is not equivalence'],
    ['Simple mechanical binding constraint','Restriction becomes consequential near/exceeding planned eligible spending','Cash advantage may increase with r under the simple extra-spending constraint account','Narrow conditional prediction, not all standard theory; r baseline association not causal'],
    ['Label / mental accounting','Level gaps may persist even when inframarginal','No unique monotone slope prediction without added assumptions','Possible interpretation; no direct process measure'],
    ['Scale-induced categorization / endogenous earmarking','Small Cash unusually spendable; larger Cash self-categorized','Cash advantage attenuates with scale','Untested interpretation absent relative-collapse/direct process evidence']]
    mt=pd.DataFrame(mech,columns=['account','level_prediction','slope_prediction','boundary'])
    mt.to_csv(OUT/'mechanism_prediction_table.csv',index=False)
    put('mechanism_prediction_table.md','# Competing conditional predictions\n\n'+md(mt)+'''\n\n**Fact:** C uses conservative monthly food-band lower bounds [0,501,1001,2001,3001,5001] multiplied by6, not invented expenditure midpoints. Zero lower bounds yield undefined r; counts in bindingness_ratio_counts.csv, not dropped silently. r bins were saved before new outcomes were analyzed; no sparse merges. Food baseline excludes daily goods and does not directly measure intended eligible expenditure, so r is a conservative proxy, not a structural constraint index. **Elimination:** results weaken a simple sufficient bindingness explanation; they do not eliminate all restrictions/label/horizon accounts. Fixed common lower6>5000 selection uses baseline only and is the same for all amount arms. All fixed-sample levels and slopes are in bindingness_ratio_tests.csv, with Holm-75 sensitivity. Level null cannot establish slope null.''')
    put('relative_scale_falsification_memo.md',f'''# ABS versus REL finite falsification

**Fact:** same unified five folds assigned once, stratified on nine randomized cells; retained across R/A/C/Q1/Q2 and mapping variants. Fold hash/counts are aggregate, no respondent assignments released. Top75 log-loss; ordinal/midpoint MSE. Models are form×ln(amount) versus form×ln(amount/income scale), not direct process measurement. Banded income mapping is NOT exact income. Two raw coding schemes have different ranges and are not naïvely combined by subtracting10. M1/M2 bounded band midpoints and open-end alternatives are in secondary.py; rank is within-scheme fractional rank and dimensionless, not monetary income.

R comparisons:\n\n{md(cv[cv['sample'].eq('R')])}

Revision criterion: at least1% loss reduction AND descriptive paired loss interval>0. None of the R focal mappings reaches materiality; none of the five samples' nine outcome/mapping comparisons improves by1%. Paired intervals condition on shared trained folds and ignore training dependence; not formal independent inferential confidence intervals. AIC/BIC and compression interactions are secondary. Positive Cash slope×log-income implies steeper compression for lower income under this sign convention; it is a sign check only and without material predictive superiority is not a mechanism pass. Full45 secondary compression tests are in relative_scale_compression.csv. **Inference:** relative-income collapse unsupported here; **Elimination:** this does not eliminate richer scale accounts. **Speculation:** scale-induced categorization remains a candidate, not Abstract/title explanation. Proposed later wording: “A scale-dependent earmarking account is one candidate explanation, but the present design does not directly test this mechanism.”
''')
    reader=pd.io.stata.StataReader(RAW); labels=reader.variable_labels(); values=reader.value_labels()
    code=[]
    for v,label in labels.items():
        role='ID' if v=='id' else ('treatment/outcome' if v.startswith(('q32','q33','q34','q35','q36','q37','q38','q39','q40','scen_')) else 'baseline/background')
        code.append(dict(variable=v,label=label,role=role,raw_value_labels=json.dumps(values.get(v,{}),ensure_ascii=False)))
    pd.DataFrame(code).to_csv(OUT/'codebook.csv',index=False)
    put('response_time_audit.md',f'''# Response-time inventory

**NOT FEASIBLE: no response-time metadata in delivered data/export.** All {len(labels)} raw fields were inspected; no start/end timestamp, duration, page timing, item timing or Qualtrics timing field is present. Platform ID, geographic codes, ordinal answers and quality filters are not timing proxies. Questionnaire text contains no timer measurement. No timing model was run. A new platform export would require author coordination; no substitute was invented.
''')
    put('questionnaire_mechanism_audit.md','''# Full questionnaire mechanism audit (read-only)

Original integrated questionnaire was read in full (questions1–31, nine randomized versions, appendix). Every respondent answers exactly one form/amount. Raw inventory is codebook.csv. Scenario question occurs afterQ1–31; Q41–49 are platform background fields. These are baseline attributes conceptually, though exact platform capture timestamps remain unknown.

| Field/domain | Exact wording / presence | Timing | Usable direct process measure? |
| --- | --- | --- | --- |
| Q7 emergency capacity | 如果您突然遇到紧急情况，需要筹集相当于您三个月收入的资金，您觉得自己能否在短期内做到？ | pre scenario | Baseline capacity only; no earmarking process |
| Q14 pressure topics | 以下哪些方面给您带来了较大的生活压力？ | pre | No: stress topics are not transfer allocation |
| Q15 future outlook | 展望未来一年，您预计自己的生活状况会如何变化？ | pre | No: outlook is not planning horizon |
| Q29 food need | 您家庭每月用于食品的支出（包括买菜、外出就餐、零食饮料等）大约是多少？ | pre | Baseline band; not intended use of transfer |
| Q30 medical need | 过去一年中，您家庭自费的医疗支出（门诊、买药、体检等，不含医保报销部分）大约是多少？ | pre | Baseline household need; not earmarking |
| Q31 prior subsidy | 过去三年，您或您的家庭是否实际收到过政府或平台发放的消费券、现金补贴等补助？ | pre | Prior experience, not process |
| Q32–40 total response | 您的【总消费】预计会比原计划多花多少钱？ | post randomized scenario | Outcome only; not separate mechanism |
| Savings/debt explanation | Cash: 基本不会额外多消费（因现金补贴节省的钱几乎全部存起来或用于还债） | outcome option | No: wording embedded in outcome, not measured amounts allocated |
| Food category1 | 基本不会额外多消费（用券购买原来计划要买的东西，省下的现金存起来或还债） | outcome option | No separate process |
| Medical category1 | 基本不会额外多消费（用补贴支付原来计划的医疗支出，省下的现金存起来或还债） | outcome option | No separate process |
| Intended use; saving/debt/housing/education/health allocation | Separate response absent | absent | NOT FEASIBLE |
| Planning/saving tendency; perceived spendability; mental account/earmarking; numeracy | Absent | absent | NOT FEASIBLE; no psychological substitutions |
| Open text; reasons | Absent | absent | NOT FEASIBLE |
| Other form/amount | Only Cash/Food/Medical and200/1000/5000 | randomized | None available; no extra arms |

**Inference:** no direct mechanism test can be proposed under D3. Stop rather than mine substitutes. No latent trait, ML/HTE, SHAP or new moderator search.
''')
    put('positive_control_literature.md','''# Positive-control direction audit

Directions were recorded before interpreting this round's control coefficients: Cash-size negative; Q7 ability/capacity negative (greater capacity, less liquidity constraint); baseline income-rank negative. This is a finite diagnostic set, not a universal theoretical law. Three outcomes in one Cash-only joint model with z, standardized capacity and income rank; Holm9.

[Jappelli–Pistaferri2014](https://doi.org/10.1257/mac.6.4.107) and [2020](https://doi.org/10.1257/pol.20180420) support reported-MPC/liquidity cash-on-hand associations, not equivalence between personal income and liquid wealth. [Parker2017](https://doi.org/10.1257/mac.20150331) provides realized household responses to a payment experiment; [Kaplan–Violante–Weidner2014](https://doi.org/10.1353/eca.2014.0002) shows wealthy hand-to-mouth households can have little liquid wealth. [Andreolli–Surico2026](https://doi.org/10.1257/mac.20220169) supports size heterogeneity contingent on affluence/liquidity, not all-person negative size law. [Fuster–Kaplan–Zafar2021](https://doi.org/10.1093/restud/rdaa076) documents hypothetical gains/losses/news/loans and can find larger-gain extensive response increases. Neither direction may be redefined after current data.

**Fact:** Top75 size passes; capacity direction is nominal, Holm .06291; income positive and Holm-significant, opposite frozen sign. Ordinal/midpoint income is also positive. **Inference:** MIXED construct-validity diagnostic, qualified because income measure is banded personal income and not a clean cash-on-hand measure. Income prediction failure cannot be dismissed as a pass. **Elimination:** these checks neither validate stated versus realized spending nor reject all survey information. See full9 tests in positive_controls.csv.
''')
    put('raking_method.md','''# Adult calibration sensitivity — NOT FEASIBLE in this run

We did not obtain a verified, adult18+ compatible joint age×education official benchmark. The NBS2020 census bulletin [official interpretation](https://www.stats.gov.cn/xxgk/jd/sjjd2020/202105/t20210513_1817408.html) gives college15,467 per100,000 among *all ages*, not adults; schooling averages15+ and working-age16–59 have other universes. Applying those percentages to adults5480 would be indefensible. The [census yearbook](https://www.stats.gov.cn/sj/pcsj/rkpc/7rp/zk/indexch.htm) contains scanned age/education tables; A0401/A0402 image retrieval via the current web reader timed out. No verified18+ marginal extraction was obtained. This does NOT mean suitable official tables do not exist.

Status: NOT FEASIBLE with verified benchmarks accessible in this run. No raking weights estimated; no weighted focal analysis performed; no nationally representative claim. Summary/result CSVs contain status, not invented numeric values. Authors can provide an authorized verified adult benchmark for a later separately specified calibration; that is outside this finite audit.
''')
    for n in ['raking_weights_summary.csv','raking_focal_results.csv']:
        pd.DataFrame([dict(status='NOT FEASIBLE',reason='No verified/retrieved compatible official18+ age-by-education margins in this run',N_adult=5480,weighted_analysis_performed=False)]).to_csv(OUT/n,index=False)
    put('author_information_needed.md','''# Author information needed — cannot infer

Questionnaire appendix attributes background fields to Tencent survey platform, but this is not a verified provider contract, sampling/recruitment log or panel specification. Intro promises a cash red packet within one week and asks no duplicate entries; this is not actual payment evidence or deduplication implementation.

Authors must supply: recruitment platform/provider and role; exact field dates; sampling frame and recruitment channels; eligibility/age policy; quotas; invitations, starts/completions and definition of completion rate; actual incentive amount/payment/compliance; technical deduplication and exclusion logs; sample-size determination and stopping rule; randomization implementation/probabilities; arm-specific starts/completions/attrition if any; consent text and acquisition method; institution/committee and approval/exemption ID/date; minors coverage/parental assent if applicable; data-sharing consent and lawful access restrictions. Structural empty eight scenario fields are randomization, not attrition. Do not infer collection dates from filenames or committee IDs from template placeholders.
''')
    put('ethics_minors_audit.md','''# Ethics and minor coverage — SUBMISSION BLOCKER

**Fact:** full5497, adult5480, 17 respondents age<18. Neither the read-only current manuscript placeholders nor available questionnaire/project documents establish committee approval/exemption number, dates, consent acquisition or coverage for those minors. The questionnaire academic/aggregate-anonymity introduction is not itself committee approval or evidence of guardian consent. No actual absence of approval is alleged: documentation is missing.

**Required author action:** supply verifiable committee/approval/exemption and consent records, including minors eligibility/guardian requirements under the applicable institution. Until clarified, recommend adult-only primary sample n=5480 and showR sensitivity. The already saved revision family keeps both samples; do not quietly delete minors to improve a result. A-only omnibus globalP=.444111 does not alter the main strength decision. Missing ethics documentation is a SUBMISSION BLOCKER even if minors are excluded. No respondent identity or raw records released.
''')
    put('open_science_release_checklist.md','''# Release checklist — no raw-data release authorized

Available: finite revision manifest/history; executable scripts; aggregate numerical source tables;67-field labeled codebook including original value-label schemes; package versions; README reproduction commands; verification and figure audit. Prior preparation/results are repository dependencies; not independent replications. All new artifacts are inside this output directory.

Before submission: authors verify provenance, ethics and consent for sharing; choose stable archive/DOI for code and aggregate source; supply approved de-identified respondent data only if permissible, otherwise a real controlled-access process and reason. Current GitHub PR is not automatically a permanent DOI. Do not state that data are openly available when no respondent data release exists. Raw dta, IDs, fold/person assignments, response rows and manuscript/consent documents are not uploaded. No OSF/Zenodo account or sensitive data publication was attempted. Figure/table source CSVs are aggregate; retain README and environment for readers.
''')
    put('NC_submission_requirements.md','''# Current Nature Communications submission audit

Checked2026-10-03 against official publisher pages (cached primary-page search text where direct pages redirected to an inaccessible identity-provider gate; forms were not completely inspected). Recheck publisher's live submission forms at actual submission.

| Requirement | Current package / author action | Primary source |
| --- | --- | --- |
| Behavioural/social sciences reporting and preregistration declaration | Declare NOT preregistered; complete human behavioural/social science reporting-summary module; report samples, exclusions, nulls and uncertainty | https://www.nature.com/ncomms/submit/hum-behav-soc-sci-studies |
| New custom code for central conclusions available to editors/reviewers | Scripts and finite manifest included; archive/permanent access still author task | https://www.nature.com/ncomms/submit/how-to-submit |
| Numerical source data behind graphs/tables | Aggregate CSVs provided; package permitted text/Excel/zip as required at upload | https://www.nature.com/ncomms/submit/how-to-submit |
| Data/Code Availability with real access restrictions | No blanket public respondent-data claim; approval/consent and access procedure required | https://www.nature.com/documents/nr-data-availability-statements-data-citations.pdf |
| Code-sharing/permanent identifiers | GitHub currently; stable DOI archive optional next author action | https://support.nature.com/en/support/solutions/articles/6000237619-software-and-code-sharing |
| Reporting/checklist resources | Complete applicable reporting checklist rather than assume audit substitutes for form | https://www.nature.com/ncomms/submit/resources |
| Authorship, consent and contributions | All authors approve submission/contributions; committee/consent details unresolved | https://www.nature.com/ncomms/editorial-policies/authorship |

Submission currently blocked by missing ethics/consent documentation. This audit does not certify journal acceptance, reviewer compliance or national representativeness. No manuscript changed.
''')
    ref=read('reference_audit.csv')
    # Compare literal citation date/pages with verified deposited metadata; no silent replacement.
    issues=[]
    for r in ref.itertuples():
        citation_year=re.findall(r'\b(?:19|20)\d{2}\b',r.citation)
        published_year=re.search(r'\d{4}',str(getattr(r,'year','')))
        issues.append(dict(ref_number=r.ref_number,citation_year=citation_year[-1] if citation_year else '',metadata_date=getattr(r,'year',''),year_matches=not published_year or not citation_year or citation_year[-1]==published_year[0],pages_match=str(r.pages).replace('-', '–') in r.citation if pd.notna(r.pages) else 'not applicable',action=r.action))
    pd.DataFrame(issues).to_csv(OUT/'reference_metadata_checks.csv',index=False)
    prior_crossfit=pd.read_csv(OUT.parent/'mpc_allx_screen/crossfit_family_summary.csv')
    prior_crossfit.to_csv(OUT/'allx_crossfit_reused_summary.csv',index=False)
    put('allx_supplemental_reuse.md','# Existing all-X cross-fit diagnostics (not rerun)\n\n'+md(prior_crossfit)+'''\n\nSource: results/mpc_allx_screen/crossfit_family_summary.csv and crossfit_stability.csv; five cell-stratified folds, seed20261005. Every eligible X×three outcomes×Cash/RC was directionally projected from training coefficients into held-out coefficients. Factor vectors use training unit direction. Training sign comparisons to overlapping full sample are optimistic; heldout same-sign counts are descriptive and dependent. Sparse factor/zero-variance failures are retained. This is not a joint BLP/GATES heterogeneity test and does not establish universality. No valid already-completed joint BLP/GATES test for this specific size-curve estimand was found in the cited all-X package; older HTE analyses address different estimands and are not recycled as validation. No new black-box/global heterogeneity search was launched. Reuse is supplemental, not independent replication.''')
    put('reference_date_note.md','''# Bibliographic date reconciliation

Publisher-deposited published dates can be online-first dates, not issue dates. Ref9 JEEA15(1),99–127 is2017 issue with2016-12-22 online record; ref20 REStud88(4),1760–1795 is2021 issue with2020-11-09 online record. Both manuscript issue-year citations are appropriate, not nonexistent papers. The year_matches=False rows in reference_metadata_checks.csv denote literal online-date differences, not a recommendation to replace issue years. All current page ranges/article IDs match deposited metadata. Reference15 remains working-paper status; no journal publication invented. Metadata and claim scopes are separate checks, and no manuscript edited.
''')
    put('NC_revision_literature_map.md','''# Problem-driven literature map

All30 current references have title/first-author checks and publisher-deposited metadata or publisher working-paper-page verification in reference_audit.csv; this is not a full-text replication of each claim. Exact authors, status, dates, pages and DOIs are retained. Compare current citation against metadata_checks.csv; page abbreviations and article IDs may require manual formatting, not scientific correction. Claim support is explicitly bounded below. Ref22 DOI lookup initially failed; AEA/NBER identify10.1257/mac.20150331, now verified. Ref15 is NBER35698, September2026, a working paper not peer-reviewed journal. Never transfer eligible-category MPC estimates to total consumption.

| Paper | Journal/year | Exact relevant result / limit | Intro paragraph | Action |
| --- | --- | --- | --- | --- |
| Raghubir–Srivastava, The Denomination Effect | JCR2009,36(4)701–713; https://doi.org/10.1086/599222 | Larger physical denomination changes spending propensity at fixed total value; NOT assigned transfer-size law | 2 | add bounded comparator |
| Andreolli–Surico, Less Is More | AEJMacro2026,18,34–68; https://doi.org/10.1257/mac.20220169 | Hypothetical size effects differ by affluence/liquidity; not universal negative slope | 2 | keep |
| Fuster–Kaplan–Zafar | REStud2021,88,1760–1795; https://doi.org/10.1093/restud/rdaa076 | Hypothetical gains/losses/news/loans; larger gains can increase extensive spending response | 2 | keep counterpoint |
| Johnson–Parker–Souleles2006; Parker et al2013 | AER; references16–17 | Realized spending after rebates/stimulus; measurement and horizon differ | 2 | keep |
| Jappelli–Pistaferri2014/2020; Parker2017 | AEJMacro/Policy; refs19/21/22 | Reported liquidity association and experiment on household spending; do not equate personal income with cash-on-hand | 2,7 | keep |
| Hastings–Shapiro | AER2018,108,3493–3540; https://doi.org/10.1257/aer.20170866 | SNAP strongly affects eligible-food spending even with inframarginality; not total-spending MPC | 3 | keep |
| Abeler–Marklein | JEEA2017,15,99–127; https://doi.org/10.1093/jeea/jvw007 | Labels alter experimental consumption despite fungibility | 3 | keep |
| Kan–Peng–Wang, Understanding Consumption Behavior: Evidence from Consumers' Reaction to Shopping Vouchers | AEJPolicy2017,9(1),137–153; https://doi.org/10.1257/pol.20130426 | Taiwan voucher survey; vendor discounts and programme context matter, not same experiment | 3 | add |
| Boehm–Fize–Jaravel | AER2025,115,1–42; https://doi.org/10.1257/aer.20240138 | Cash-like versus expiring-card realized response differs; expiry bundled | 3 | keep |
| Bonomo–Ruffini–Schanzenbach | NBERWP35698 September2026; https://doi.org/10.3386/w35698 | Food-store MPC for one-time/monthly food .18/.38 vs cash .06/.20; not total consumption, preliminary | 3 | keep exact WP status |
| Thaler1985/1999; Shefrin–Thaler1988 | refs1–3 | Mental-accounting/behavioural-life-cycle concepts, not unique slope prediction | 4 | keep concepts |
| Milkman–Beshears | JEBO2009,71,384–394; https://doi.org/10.1016/j.jebo.2009.04.007 | Small windfalls and grocery basket/hedonic purchases | 4 | keep bounded |
| Boutros | JFE2026,176,104174; https://doi.org/10.1016/j.jfineco.2025.104174 | Finite planning horizon model is candidate, not tested current process | 4,7 | keep Discussion only |
| Bernard | BundesbankDP13/2023; ref25 publisher link | Mental-account MPC framing, no direct process measurement in our survey | 4 | keep WP status |
| Lewis–Melcangi–Pilossoph | REStud2026,93,3274–3302; https://doi.org/10.1093/restud/rdaf102 | Latent realized-MPC variation weakly explained by observables; not universal current effect | 7 | keep bounded |
| Kaplan–Violante–Weidner | BPEA2014,77–138; https://doi.org/10.1353/eca.2014.0002 | Wealthy hand-to-mouth implies income is not liquidity | 2,7 | keep |
| Hahnel energy; general value-context2026; neural-context2016 | refs4–6 | Different decision/process outcomes; no transfer-size MPC evidence | none | remove from core Intro unless specific relevance is established |
| Gennetian2024; Magnuson2025; McGuire2022 | refs28–30 Nature-family | Family investment/process/wellbeing, not form×size MPC | none | remove brand decoration from core Intro |

This is a suggested citation map only. No manuscript references or prose were edited. Existing30 references retained in audit with bounded claims and keep/remove actions, even if not individually listed above.
''')
    put('revised_intro_outline.md','''# Revised Introduction outline only

1. Broad problem: transfer-size/form interactions matter for whether a response at one amount extrapolates; distinguish stated responses from realized policy effects.
2. Classic size/MPC evidence: stimulus/rebate responses, reported versus realized measurement, liquidity and magnitude; explicitly note heterogeneous/opposite size patterns.
3. Form/fungibility/labeling: SNAP, labels, coupons/Taiwan, cash-versus-expiring card; distinguish eligible-category from total spending, bundled attributes from pure restrictions.
4. Competing predictions: full fungibility, conditional mechanical bindingness, label levels without unique slopes, scale-categorization as untested candidate.
5. Design: one randomized scenario per respondent, three forms×three amounts, six ordered hypothetical total-response bins; no repeated individual curves.
6. Empirical pattern: strongest observed Cash upper-tail decline, separately displayed Food/Medical; finite-family global correction fails, so suggestive not a strong headline.
7. Bounded contribution: context-specific evidence relevant to extrapolation, no identified earmarking mechanism, underpowered moderation, mixed measurement controls and external-validity limits.

Outline is an audit asset, not rewritten manuscript paragraphs.
''')
    put('manuscript_hygiene_assets.md','''# Later manuscript revision checklist — no edits performed

Use actual submitted figures rather than placeholders. Show Food/Medical separately before equal-weight derived Restricted. Remove PR14/13, approved, reviewer-style critique, revised narrative structure and branch/commit/workflow wording from future manuscript. Never imply preregistration. Avoid repeated numbers across Abstract/Intro/Results/Discussion; use one main factual summary and tables for full uncertainty. Global correction fails: avoid “reshapes”, strong causal-sounding mechanism titles and “universal” from moderator nulls. Use suggests/evidence for and explicitly describe revision-stage analysis history. Gap attenuation is not equality; no equivalence without margin limitations. “A transfer-form effect estimated at one amount need not extrapolate to another amount” is a bounded implication, not a demonstrated universal law. Do not use mental-account mechanism as Abstract explanation without direct evidence. Ethics/recruitment/data access documentation must be supplied before submission.
''')
    put('figure_captions.md','''# Figure legends and aggregate sources

All forms are randomized scenario assignments; outcomes hypothetical stated responses. Arial, white background, common form palette, editable vector PDF/SVG and600dpi PNG. Points/CIs are observed/model summaries, no smoothing or continuous dose-response claims.

1. **fig_form_top75:** Cash, Food, Medical category6 shares at ¥200/1000/5000. R, n=5497; normal95% cell-mean intervals. Source form_amount_profiles.csv, outcome=top75. Straight segments guide discrete comparisons only; no fitted continuous curve.
2. **fig_form_extended:** separate-form ordinal1–6 and midpoint-coded [.0,.05,.175,.375,.625,.875] responses. Same cells/n/interval method; source form_amount_profiles.csv. Coding does not reveal true within-bin MPC.
3. **fig_specification_family:** all1,500 tests' scalar/vector coefficients and pointwise95% intervals, separate outcomes/model/amount units; saturated tests show both components. Omnibus components overlap primary decomposition by design. Source specification_figure_source.csv and specification_grid.csv; no significance sorting or selected winner. No scalar omnibus estimate exists.
4. **fig_quality_gradient:** Cash−Food and Cash−Medical Top75 slopes in R/A/C/Q1/Q2, HC3 normal95% intervals. Fivefold trend units, common vertical scale. Source quality_gradient_source.csv. Quality rules change composition and power; not an exogenous quality intervention.
5. **fig_bindingness_ratio:** R Top75 Cash−Food mean gaps by five saved ratio bins, HC3 normal95% intervals. Source bindingness_ratio_source.csv; ratio uses amount/conservative6monthfood lower bound; undefined counts separately. Unadjusted baseline-dependent associations, not causal dose effects.
6. **fig_relative_scale:** held-out relative loss improvement of income-normalized versus absolute-dose model, M1/M2/rank, three outcomes, five samples, unified5fold. Source relative_scale_cv.csv. Dotted1% rule is revision-stage materiality criterion, not preregistration; no independent cross-validation inference is claimed.
7. **fig_allx_supplemental:** raw-P QQ plots for six original primary cash/rc×ordinal/midpoint/top75 families. Source allx_QQ_source.csv; complete q distribution/counts in allx_q_distribution.csv/allx_supplemental_counts.csv. No nominal winner; raw/constructed representations are dependent. Reused screen, not a new independent sample.
''')
    versions={x:importlib.metadata.version(x) for x in ['numpy','pandas','scipy','statsmodels','scikit-learn','matplotlib','python-docx','pypdf','Pillow']}
    put('requirements.txt','\n'.join(f'{k}=={v}' for k,v in versions.items()))
    inputs=[RAW,QUESTION,MAN,OUT/'revision_spec_manifest.json',OUT.parent/'mpc_allx_screen/all_screen_results.csv',OUT.parent/'mpc_size_curve/size_curve.py',OUT.parent/'mpc_final_strengthening/strengthening.py']
    (OUT/'input_checksums.json').write_text(json.dumps({str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    put('README.md','''# Reproduce NC revision audit

Use Python3.11+ and requirements.txt; Windows Arial/Poppler for visual QA. Run from repository root. Raw data and questionnaire/manuscript stay local and require author access, not bundled. Replace local paths in report_package.py/reference_audit.py for document audit; analysis accepts --data. Historical preparation/AME/screen dependencies are inherited in this branch. No previously stored respondent data is loaded/exported by the new reporting scripts.

```text
python 消费调查/results/nc_revision_audit/run_revision.py --data <authorized_raw.dta>
python 消费调查/results/nc_revision_audit/secondary.py --data <authorized_raw.dta>
python 消费调查/results/nc_revision_audit/reference_audit.py
python 消费调查/results/nc_revision_audit/figures.py
python 消费调查/results/nc_revision_audit/report_package.py
python 消费调查/results/nc_revision_audit/verify_package.py --data <authorized_raw.dta>
```

Reference audit makes network requests to publisher-deposited Crossref metadata (proxy configuration in script may be removed for another machine), with bounded claim scopes; reference metadata can change. Online NC/census audit is dated2026-10-03. Seed/fold/multiplicity settings are in manifest and scripts. Do not edit the saved manifest after examining results. All numerical exports are aggregate. qa/ renders are ignored; no IDs/raw rows/fold assignment release. NOT FEASIBLE and unresolved author items are not passed tests. See NC_REVISION_RESULT.md for finalB decision. No further exploratory work follows.
''')
    put('RESULT.md','''# Work result

## Summary

Executed NC_REVISION_WORK finite audit. Recommended strengthB: suggestive; focal global correctionNO. Manuscript untouched; no unplanned analysis/raw-data release. Output lives only in this directory.

## Files changed

New revision manifest, four analysis/audit scripts plus figure/report/verification code, aggregate model/grid/focal/bootstrap/mechanism/CV/power/control/source tables,30reference audit, seven PDF/SVG/600dpiPNG figures, decision/method/ethics/submission/reproduction reports and checksums. See directory inventory and README.

## Key implementation decisions

Global joint constrained-null score multiplier (5,000 draws) with multidimensional Wald norms, not draw-level model refits; explicitly bounded null simulation. Distinct raw income schemes; identical5fold assignments; fixed ratio bins; separateFood/Medical; Restricted exactly equal-weight derivative. Optional Bayesian analysis skipped. Timing/process and adult-calibration checksNOTFEASIBLE; ethics documentationSUBMISSIONBLOCKER.

## Testing performed

See verification_results.json for actual assertions and independent estimator checks; figure_audit.md for rendered visual inspection. Never interpretNOTFEASIBLEas passed. Code execution logs/counts in execution_manifest.json/model_diagnostics.csv. No manuscript tests or manuscript changes.

## Acceptance Criteria

Required finite outputs produced; complete estimable1,500-row family and5,000jointdraws; all sample/form displays; mechanism/MDE/positivecontrols and literature/submission audits; missingvariables honestly flagged; no expansion/manuscript edit. Publication metadata recorded below when repository push/PR completes.

## Known issues

Asymptotic score-level global correction and one pooled-null simulation, not exact/refit procedure. PairedCV intervals ignore training dependence. No verifiedadultcalibration. Missing recruitment/ethics/minor records block submission. Historical subgroup observations reuse this dataset; no independentreplication. Public respondentdata authorization unresolved.

## Branch

feature/nc-revision-audit; stacked on feature/mpc-hero-figure (prior#14) plus latestmain task plan.

## Commit SHA

Pending publication.

## PR

Pending publication; do not merge automatically.
''')
    print('Reports, codebook and input checksums generated; manuscript read-only.')
    if (OUT/'publication.json').exists():
        pub=json.loads((OUT/'publication.json').read_text(encoding='utf-8'))
        p=OUT/'RESULT.md';s=p.read_text(encoding='utf-8')
        s=s.replace('Pending publication.',pub['analysis_commit'])
        s=s.replace('Pending publication; do not merge automatically.',pub['pr_url']+'; stacked on '+pub['base']+'; not merged.')
        p.write_text(s,encoding='utf-8')
if __name__=='__main__':main()
