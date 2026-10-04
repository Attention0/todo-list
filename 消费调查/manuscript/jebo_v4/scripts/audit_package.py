"""Source, claim-coverage and document checks; never estimates a new model."""
from pathlib import Path
import csv,json,re,hashlib,subprocess,zipfile
from docx import Document
from docx.oxml.ns import qn
M=Path(__file__).resolve().parents[1];R=M.parents[1];O=R/'results/jebo_review_revision';REPO=R.parent
BASE='363313c8adc52cce252c11390a32e46c7a98d25f'
def read(fn):return list(csv.DictReader((O/fn).open(encoding='utf8')))
s=(M/'JEBO_manuscript_v4.md').read_text(encoding='utf8');body=s.split('## References')[0]
checks=[]
def check(name,passed,detail=''):
    checks.append(dict(check=name,passed=bool(passed),detail=detail))
    if not passed:print('FAIL',name,detail)
ab=body.split('## Abstract')[1].split('Keywords:')[0];intro=body.split('## 1 Introduction')[1].split('## 2 Conceptual')[0]
check('Abstract length',170<=len(ab.split())<=220,str(len(ab.split())))
check('Introduction length',1500<=len(intro.split())<=1900,str(len(intro.split())))
check('First three paragraphs contain puzzles',all(x in '\n'.join(intro.strip().split('\n\n')[:3]) for x in ['25.30%','22.70%','17.73%','gap']))
check('Next two paragraphs contain scale and anatomy',all(x in '\n'.join(intro.strip().split('\n\n')[3:5]) for x in ['yuan','highest','offsetting']))
check('No internal workflow vocabulary',not re.search(r'\b(PR19|PR #|GitHub|SPEC|WORK\.md|worktree)\b',body))
check('No source placeholders','{{' not in s)
check('Main figure first-mention order',[m.group(1) for m in re.finditer(r'Figure ([1-4])',body)][:3]==['1','2','3'])
high=[b[2:] for b in (M/'JEBO_Highlights_v4.md').read_text().split('\n\n') if b.startswith('- ')]
check('Five short highlights',len(high)==5 and all(len(x.strip())<=85 for x in high),str([len(x.strip()) for x in high]))
for fn in ['adult_cell_distribution.csv','adult_share_yuan_table.csv','affine_midpoint_fit.csv','incremental_yuan_response.csv','food_bindingness_raw_cells.csv','scientific_family42.csv']:
    check('Exact aggregate copy '+fn,(M/'source_data'/fn).read_bytes()==(O/fn).read_bytes())
cell=read('adult_share_yuan_table.csv');lookup={(r['form'],int(r['amount'])):r for r in cell}
for r in cell:
    check('Displayed share '+r['form']+' '+r['amount'],f"{float(r['midpoint'])*100:.2f}%" in s)
    check('Displayed yuan '+r['form']+' '+r['amount'],f"{float(r['yuan']):,.2f}" in s)
    check('Six counts reconcile '+r['form']+' '+r['amount'],sum(int(r[f'count{k}']) for k in range(1,7))==int(r['N']))
changes=list(csv.DictReader((M/'source_data/endpoint_category_changes.csv').open()))
for r in changes:
    expected=(float(lookup[r['form'],5000]['share'+r['category']])-float(lookup[r['form'],200]['share'+r['category']]))*100
    check('Descriptive subtraction '+r['form']+' '+r['category'],abs(float(r['change_pp'])-expected)<1e-12)
family=read('scientific_family42.csv');check('Complete scientific family',len(family)==42)
for outcome,test in [('midpoint','cash'),('top75','cash'),('ordinal','cash'),('top75','omnibus'),('top75','cash-food'),('top75','cash-medical')]:
    r=next(x for x in family if x['outcome']==outcome and x['test']==test)
    expected=(outcome in ['midpoint','top75'] and test=='cash')
    check('Frozen inference '+outcome+' '+test,(float(r['holm42'])<.05)==expected and (float(r['minP42'])<.05)==expected)
table3=s.split('| Outcome and test |')[1].split('\n\n')[0].splitlines()[2:]
for line,(outcome,test) in zip(table3,[('midpoint','cash'),('top75','cash'),('ordinal','cash'),('top75','omnibus'),('top75','cash-food'),('top75','cash-medical')]):
    vals=line.strip('| ').split('|');r=next(x for x in family if x['outcome']==outcome and x['test']==test)
    for printed,col in zip(vals[1:],['p','holm42','minP42']):
        check('Table 3 transcription '+outcome+'/'+test+'/'+col,abs(float(printed)-float(r[col]))<=max(1e-7,abs(float(r[col]))*.005))
for form,cat in [('cash',1),('cash',6),('cash',3),('food',1),('food',6),('food',3),('food',4),('medical',1),('medical',3),('medical',5),('medical',6)]:
    for amount in [200,5000]:
        value=float(lookup[form,amount]['share'+str(cat)])*100
        check(f'Anatomy number {form}/{cat}/{amount}',f'{value:.2f}%' in s)
check('Food sign preserved','contradicts the sharp prediction' in s and '+8.00' in s and '−6.67' in s)
check('Measurement boundaries',all(x in s for x in ['not observed spending participation','not structural MPC','not an exact randomization test','do not establish equivalence','unmeasured candidate construct']))
supp=(M/'JEBO_supplement_v4.md').read_text(encoding='utf8')
check('Supplement reorganized',supp.index('## S2 Original')<supp.index('## S3 Multiplicity')<supp.index('## S4 Yuan'))
check('Supplement coverage',all(x in supp for x in ['Complete 42 test','Earlier 21','Original adult six-bin','Interval affine','Historical broad observable','Q1 requires','Original questionnaire','Reference verification']))
old=(M.parent/'jebo_v3/JEBO_supplement_v3.md').read_text(encoding='utf8')
oldtables=[b.strip() for b in old.split('\n\n') if b.strip().startswith('|')]
check('All original supplementary table values retained',all(t in supp for t in oldtables),str(len(oldtables)))
oldq=old.split('## S9 Original')[1].split('## S10 ')[0]
newq=supp.split('## S9 Original')[1].split('## S10 ')[0]
check('Questionnaire Chinese and English unchanged',oldq.strip()==newq.strip())
audit=list(csv.DictReader((M/'JEBO_V4_REFERENCE_AUDIT.csv').open(encoding='utf8')))
check('All core references audited',len(audit)==25 and all(r['primary_url'] and r['claim_boundary'] for r in audit))
check('Recent statuses bounded',all(r['publication_status']=='Working paper' for r in audit if r['id'] in ['N24','N25','N30']))
check('Published Jappelli status',any(r['id']=='JS2026' and r['volume']=='186' and r['pages']=='105303' for r in audit))
for stem in ['JEBO_manuscript_v4','JEBO_supplement_v4','JEBO_Highlights_v4']:
    d=Document(M/(stem+'.docx'))
    check('Body style '+stem,d.styles['Normal'].font.name=='Times New Roman' and d.styles['Normal'].font.size.pt==12)
    check('Title style '+stem,d.styles['Title'].font.name=='Times New Roman' and d.styles['Title'].font.size.pt==17)
    check('Indent '+stem,abs(d.styles['Normal'].paragraph_format.first_line_indent.cm-.74)<.01)
    prose=[p for p in d.paragraphs if p.text.startswith(('Table 2 makes','Table 3 separates','Figure 2 displays','Figure 4 shows'))]
    check('Figure/table references retain body size '+stem,all(all(r.font.size is None or r.font.size.pt==12 for r in p.runs) for p in prose))
    with zipfile.ZipFile(M/(stem+'.docx')) as z:
        xml=z.read('word/document.xml').decode();style=z.read('word/styles.xml').decode()
        check('No comments/tracked changes '+stem,not any('comments' in n for n in z.namelist()) and not re.search(r'<w:(ins|del|moveFrom|moveTo)\b',xml))
        check('Page field '+stem,any(b'PAGE' in z.read(n) for n in z.namelist() if re.match(r'word/footer\d+\.xml$',n)))
    check('Images '+stem,len(d.inline_shapes)=={'JEBO_manuscript_v4':4,'JEBO_supplement_v4':1,'JEBO_Highlights_v4':0}[stem])
changed=subprocess.check_output(['git','diff','--name-only','--diff-filter=MDR',BASE],cwd=REPO,text=True,encoding='utf8').strip()
check('No tracked PR19 file modified or deleted',not changed,changed)
branch=subprocess.check_output(['git','branch','--show-current'],cwd=REPO,text=True).strip()
check('Requested branch',branch=='feature/jebo-v4-narrative-rewrite')
for fn in ['JEBO_V4_WRITING_SPEC.md','JEBO_V4_WORK.md']:
    wanted=subprocess.check_output(['git','show','origin/main:消费调查/'+fn],cwd=REPO)
    check('Main authority '+fn,(R/fn).read_bytes().replace(b'\r\n',b'\n')==wanted.replace(b'\r\n',b'\n'))

# Paragraph-level ledger is intentionally complete (including interpretation and limitations).
# Multiple tags reflect mixed empirical/interpretive paragraphs; source pointers preserve context.
mapping={
'Abstract':('randomized cell fact;coded/descriptive transformation;inferential result;interpretation;limitation','adult_share_yuan_table.csv;adult_cell_distribution.csv;scientific_family42.csv;food_bindingness_sign_reproduction.csv'),
'1':('randomized cell fact;coded/descriptive transformation;literature claim;interpretation;limitation','adult_share_yuan_table.csv;adult_cell_distribution.csv;scientific_family42.csv;JEBO_V4_REFERENCE_AUDIT.csv'),
'2':('literature claim;interpretation;limitation','JEBO_V4_REFERENCE_AUDIT.csv;QUESTION_INTERVAL_AUDIT.md'),
'3':('randomized cell fact;inferential result;limitation','REVIEWER_REANALYSIS_MANIFEST.md;sample_characteristics.csv;randomization_balance.csv;QUESTIONNAIRE_APPENDIX.md;QUESTION_INTERVAL_AUDIT.md'),
'4':('randomized cell fact;coded/descriptive transformation;inferential result;limitation','adult_share_yuan_table.csv;incremental_yuan_response.csv;endpoint_elasticity.csv;scientific_family42.csv;family42_simultaneous_intervals.csv'),
'5':('randomized cell fact;coded/descriptive transformation;interpretation;limitation','adult_cell_distribution.csv;bottom_category_decomposition.csv;source_data/endpoint_category_changes.csv'),
'6.1':('mechanism diagnostic;literature claim;interpretation;limitation','scientific_family42.csv;JEBO_V4_REFERENCE_AUDIT.csv'),
'6.2':('coded/descriptive transformation;inferential result;limitation','affine_midpoint_parameters.csv;affine_model_comparison.csv;interval_affine_diagnostics.csv;interval_affine_fit.csv'),
'6.3':('coded/descriptive transformation;mechanism diagnostic;limitation','adult_cell_distribution.csv;bottom_category_decomposition.csv;QUESTION_INTERVAL_AUDIT.md'),
'6.4':('mechanism diagnostic;inferential result;limitation','food_bindingness_sign_reproduction.csv;food_bindingness_ratio.csv;food_bindingness_raw_cells.csv;FOOD_BINDINGNESS_SIGN_AUDIT.md'),
'6.5':('mechanism diagnostic;inferential result;limitation','food_bindingness_continuous.csv;food_bindingness_adjusted.csv;MECHANISM_PREDICTION_PRECISION_LEDGER.csv'),
'6.6':('mechanism diagnostic;interpretation;limitation','MECHANISM_PREDICTION_PRECISION_LEDGER.csv;../nc_evidence_revision/mechanism_tests.csv;JEBO_supplement_v4.md S7'),
'6.7':('mechanism diagnostic;interpretation;limitation','../nc_revision_audit/allx_supplemental_counts.csv;../nc_revision_audit/allx_crossfit_reused_summary.csv;MECHANISM_PREDICTION_PRECISION_LEDGER.csv;JEBO_supplement_v4.md S8'),
'7':('interpretation;limitation','V4_STORY_MEMO.md;adult_cell_distribution.csv;QUESTIONNAIRE_APPENDIX.md;JEBO_V4_REFERENCE_AUDIT.csv'),
'8':('interpretation;limitation;inferential result','scientific_family42.csv;adult_cell_distribution.csv;MECHANISM_PREDICTION_PRECISION_LEDGER.csv;STUDY_METADATA_AUDIT.md')}
ledger=[];section='Abstract';pid=0
for block in body.split('\n\n'):
    block=block.strip()
    if block.startswith(('## ','### ')):section=block.lstrip('# ');continue
    if not block or block.startswith(('# ','![','Keywords:','JEL classification:')):continue
    key=section.split()[0];key=key if key in mapping else key.split('.')[0]
    tags,source=mapping.get(key,('limitation','AUTHOR_INFORMATION_REQUIRED.md'))
    pid+=1;flags=[]
    for name,pattern in [('spendability','spendability'),('equivalence','equivalen'),('bottom margin','participation|above-bottom'),('moderation','moderat|homogene'),('interaction','interaction'),('interval structural','structural|latent'),('distribution convergence','converg|identical')]:
        if re.search(pattern,block,re.I):flags.append(name)
    ledger.append(dict(claim_id=f'V4-P{pid:03d}',section=section,claim_type=tags,claim=block,source=source,review_flags=';'.join(flags),review_disposition='Editorial review: descriptive scope / explicit uncertainty / candidate language retained; no new test.'))
with (M/'JEBO_V4_CLAIM_LEDGER.csv').open('w',newline='',encoding='utf8') as f:
    w=csv.DictWriter(f,list(ledger[0]));w.writeheader();w.writerows(ledger)
expected_blocks=[b for b in body.split('\n\n') if b.strip() and not b.strip().startswith(('#','![','Keywords:','JEL classification:'))]
check('Complete paragraph claim ledger',len(ledger)==len(expected_blocks),str(len(ledger)))
report=dict(base=BASE,branch=branch,checks=checks,passed=sum(x['passed'] for x in checks),total=len(checks),abstract_words=len(ab.split()),intro_words=len(intro.split()),main_words_before_references=len(body.split()),claim_paragraphs=len(ledger),new_inferential_models=0)
(M/'verification.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps({k:v for k,v in report.items() if k!='checks'},indent=2))
raise SystemExit(0 if all(x['passed'] for x in checks) else 1)
