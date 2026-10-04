"""Render editable manuscripts from reviewed Markdown; no empirical analysis."""
from pathlib import Path
import re,json,sys
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
M=Path(__file__).resolve().parents[1]
def newdoc():
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2.1)
    for st in d.styles:
        if st.type==1:
            st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
            fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
            if fonts is not None:
                for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
                    fonts.attrib.pop(qn('w:'+attr),None)
                for attr in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
                fonts.set(qn('w:eastAsia'),'SimSun')
    p=d.styles['Normal'];p.font.size=Pt(12);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6);p.paragraph_format.first_line_indent=Cm(.74);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.widow_control=True
    for name,size in [('Title',17),('Heading 1',14),('Heading 2',12),('Heading 3',12)]:
        s=d.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.1;s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Title'].paragraph_format.space_before=Pt(0)
    p=d.sections[0].footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.subject='Stated spending responses across transfer forms and amounts'
    return d
def addtext(p,t):
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',t)):
        r=p.add_run(part);r.bold=bool(i%2)
def tab(d,b):
    rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in b.splitlines() if not re.match(r'^\|[\s:|\-]+\|$',line)]
    n=len(rows[0]);assert all(len(r)==n for r in rows)
    if n>9:
        # Keep every source value, divide wide output into readable column panels.
        for part,ix in enumerate([list(range(6)),[0]+list(range(6,n))],1):
            p=d.add_paragraph('Column panel '+str(part));p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.size=Pt(9)
            rr=[[row[k] for k in ix] for row in rows]
            tab(d,'\n'.join('| '+' | '.join(row)+' |' for row in rr))
        return
    t=d.add_table(rows=0,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    total=6.61
    if n==3:w=[1.34,2.43,2.84]
    elif n==4:w=[2.53,1.36,1.36,1.36] if 'test' in rows[0][0] else [1.45,1.72,1.72,1.72]
    elif rows[0][0]=='Bundle':w=[.67,2.64,1.1,1.1,1.1]
    elif rows[0][0]=='Form' and n==9:w=[.58,.53,.42]+[.8467]*6
    elif n==5:w=[1.6]+[1.2525]*4
    else:
        first=1.16 if n>=7 else 1.55;w=[first]+[(total-first)/(n-1)]*(n-1)
    if rows[0][0]=='account':w=[1.66,.4]+[(total-2.06)/(n-2)]*(n-2)
    if rows[0][:3]==['variable','category','label']:w=[.93,.72,1.79]+[(total-3.44)/(n-3)]*(n-3)
    if rows[0][0]=='comparison':w=[2.35]+[(total-2.35)/(n-1)]*(n-1)
    if rows[0][0]=='model' and n>=7:w=[1.2,1.18]+[(total-2.38)/(n-2)]*(n-2)
    if n>=10:w=[total/n]*n
    if rows[0][0]=='Target' and n==6:w=[.52,.7,.95,1.22,1.61,1.61]
    if rows[0][0]=='Work':w=[1.4,2.1,3.11]
    for c,v in zip(t.columns,w):c.width=Inches(v)
    for j,row in enumerate(rows):
        cells=t.add_row().cells;pr=cells[0]._tc.getparent().get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if j==0:pr.append(OxmlElement('w:tblHeader'))
        for i,(c,v) in enumerate(zip(cells,row)):
            c.width=Inches(w[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            aliases={'above_bottom_change':'Above-bottom change','conditional_above_bottom_change':'Conditional code change','bottom_category_component':'Bottom component','conditional_above_bottom_component':'Conditional component','bottom_lo':'Bottom low','bottom_hi':'Bottom high','conditional_lo':'Conditional low','conditional_hi':'Conditional high','MDE80_nominal':'MDE80 nominal','MDE80_family':'MDE80 family','moderator_p10_p90_span':'10th–90th span','compatible_change_lo':'Change low','compatible_change_hi':'Change high','change_lo_over_absolute_headline':'Low / headline','change_hi_over_absolute_headline':'High / headline','max_abs_target_discrepancy':'Max. target gap','relative_gain_vs_additive':'Relative gain vs additive','absolute_gain_vs_additive':'Absolute gain vs additive','relative_gain_vs_ABS':'Relative gain vs ABS','absolute_gain_vs_ABS':'Absolute gain vs ABS'}
            if j==0:c.text=aliases.get(v,v.replace('_',' '))
            else:c.text=('above-bottom' if v=='any_spending' else v.replace('_',' '))
            props=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
            props.append(mar)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D0D0D0');borders.append(e)
            props.append(borders)
            if j==0:e=OxmlElement('w:shd');e.set(qn('w:fill'),'EEEEEE');props.append(e)
            for p in c.paragraphs:
                f=p.paragraph_format;f.first_line_indent=Cm(0);f.line_spacing=1.05;f.space_after=Pt(2);f.space_before=Pt(2);f.keep_with_next=j==0 or (len(rows)<=10 and j<len(rows)-1)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if (i==0 or n==3 or len(v)>24) else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9);r.bold=j==0
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.paragraph_format.first_line_indent=Cm(0)
    for r in p.runs:r.font.size=Pt(2)
def build(stem):
    d=newdoc();s=(M/(stem+'.md')).read_text(encoding='utf8');refs=False
    for b in s.split('\n\n'):
        b=b.strip()
        if not b:continue
        if b.startswith('# '):d.add_paragraph(b[2:],'Title');d.core_properties.title=b[2:]
        elif b.startswith('## '):
            refs=b=='## References';p=d.add_paragraph(b[3:],'Heading 1')
            if stem=='JEBO_supplement_v4_1' and not b.startswith('## S1 '):p.paragraph_format.page_break_before=True
        elif b.startswith('### '):d.add_paragraph(b[4:],'Heading 2')
        elif b.startswith('|'):tab(d,b)
        elif b.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',b);p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
            r=p.add_run();r.add_picture(str(M/m.group(2)),width=Inches(6.6));r._r.xpath('.//wp:docPr')[0].set('descr',m.group(1))
        else:
            if b.startswith('- '):
                p=d.add_paragraph('• '+b[2:]);p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;continue
            p=d.add_paragraph();addtext(p,b)
            if re.match(r'^(Table|Figure) S?\d+[a-z]?\.',b) or b.startswith(('Note:','Keywords:','JEL classification:')):
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(10)
                if re.match(r'^Table S?\d+[a-z]?\.',b):
                    p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.bold=True
            if refs:
                p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5);p.paragraph_format.line_spacing=1.05;p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.size=Pt(10)
    for x in list(d.styles.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
    d.save(M/(stem+'.docx'))
    return {'file':stem+'.docx','paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes)}
if __name__=='__main__':
    result=[build(s) for s in (sys.argv[1:] or ['JEBO_manuscript_v4_1','JEBO_supplement_v4_1','JEBO_Highlights_v4_1'])]
    (M/'docx_structure.json').write_text(json.dumps(result,indent=2));print(result)
